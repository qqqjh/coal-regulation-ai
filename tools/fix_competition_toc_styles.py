from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt


DOCX = Path("D:/work/project/coal-regulation-ai") / "\u7814\u7535\u8d5b\u6280\u672f\u8bba\u6587_\u538b\u7f29\u683c\u5f0f\u7248.docx"


def set_font(run, size=10.5, bold=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "\u5b8b\u4f53")


def main() -> None:
    doc = Document(str(DOCX))
    toc_start = None
    actual_first_chapter = None
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip() == "\u76ee  \u5f55":
            toc_start = i
            continue
        if toc_start is not None and i > toc_start and p.text.strip() == "\u7b2c\u4e00\u7ae0 \u7814\u7a76\u80cc\u666f\u548c\u610f\u4e49":
            if actual_first_chapter is None:
                # first occurrence after TOC title is the TOC line
                pass
            else:
                actual_first_chapter = i
                break
            actual_first_chapter = -1
    # Actual first chapter is the second matching line after TOC.
    seen_first_chapter = 0
    for i, p in enumerate(doc.paragraphs):
        if toc_start is not None and i > toc_start and p.text.strip() == "\u7b2c\u4e00\u7ae0 \u7814\u7a76\u80cc\u666f\u548c\u610f\u4e49":
            seen_first_chapter += 1
            if seen_first_chapter == 2:
                actual_first_chapter = i
                break

    if toc_start is not None and actual_first_chapter is not None:
        for p in doc.paragraphs[toc_start + 1 : actual_first_chapter]:
            p.style = doc.styles["Normal"]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.first_line_indent = None
            p.paragraph_format.left_indent = Pt(0)
            p.paragraph_format.line_spacing = Pt(20)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            if p.text.strip().startswith(tuple(str(n) + "." for n in range(1, 8))):
                p.paragraph_format.left_indent = Pt(21)
            for r in p.runs:
                set_font(r, 10.5, False)

    for p in doc.paragraphs:
        if p.text.strip() in {"\u6458  \u8981", "Abstract", "\u76ee  \u5f55"}:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                set_font(r, 12, True)

    doc.save(str(DOCX))


if __name__ == "__main__":
    main()
