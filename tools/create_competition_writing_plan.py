from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "研电赛论文三人分工与项目完善建议_v1.docx"

BLUE = "1F4E79"
LIGHT_BLUE = "D9EAF7"
LIGHT_GRAY = "F2F2F2"
PALE_RED = "FCE4D6"
PALE_GREEN = "E2F0D9"
WHITE = "FFFFFF"
TEXT = "222222"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        tag = "w:" + edge
        node = tc_mar.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_cell_text(cell, text, bold=False, color=TEXT, size=10.5, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(str(text))
    set_run_font(run, "宋体", size, bold=bold, color=color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)


def set_run_font(run, east_asia="宋体", size=12, bold=False, color=TEXT, italic=False):
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def add_field(paragraph, field_code):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = field_code
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    normal.paragraph_format.first_line_indent = Cm(0.74)
    normal.paragraph_format.space_after = Pt(0)

    h1 = doc.styles["Heading 1"]
    h1.font.name = "Times New Roman"
    h1._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    h1.font.size = Pt(18)
    h1.font.bold = True
    h1.font.color.rgb = RGBColor.from_string(TEXT)
    h1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    h1.paragraph_format.space_before = Pt(12)
    h1.paragraph_format.space_after = Pt(8)
    h1.paragraph_format.line_spacing = 1.5
    h1.paragraph_format.first_line_indent = Cm(0)

    h2 = doc.styles["Heading 2"]
    h2.font.name = "Times New Roman"
    h2._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    h2.font.size = Pt(15)
    h2.font.bold = True
    h2.font.color.rgb = RGBColor.from_string(TEXT)
    h2.paragraph_format.space_before = Pt(8)
    h2.paragraph_format.space_after = Pt(4)
    h2.paragraph_format.line_spacing = 1.5
    h2.paragraph_format.first_line_indent = Cm(0)

    h3 = doc.styles["Heading 3"]
    h3.font.name = "Times New Roman"
    h3._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    h3.font.size = Pt(12)
    h3.font.bold = True
    h3.font.color.rgb = RGBColor.from_string(TEXT)
    h3.paragraph_format.space_before = Pt(6)
    h3.paragraph_format.space_after = Pt(2)
    h3.paragraph_format.line_spacing = 1.5
    h3.paragraph_format.first_line_indent = Cm(0)


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_run_font(r1, "宋体", 12, bold=True)
        r2 = p.add_run(text[len(bold_lead):])
        set_run_font(r2, "宋体", 12)
    else:
        run = p.add_run(text)
        set_run_font(run, "宋体", 12)
    return p


def add_compact_item(doc, label, text, color=TEXT):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.4)
    p.paragraph_format.first_line_indent = Cm(-0.4)
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.space_after = Pt(3)
    r1 = p.add_run(label)
    set_run_font(r1, "黑体", 11, bold=True, color=color)
    r2 = p.add_run(text)
    set_run_font(r2, "宋体", 11, color=color)
    return p


def add_callout(doc, title, lines, fill=LIGHT_BLUE):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Cm(15.8)
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(title)
    set_run_font(r, "黑体", 11, bold=True, color=BLUE)
    for line in lines:
        p = cell.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.35)
        p.paragraph_format.first_line_indent = Cm(-0.35)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        r = p.add_run("• " + line)
        set_run_font(r, "宋体", 10.5)
    set_cell_margins(cell, top=140, start=180, bottom=140, end=180)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_table(doc, headers, rows, widths_cm, header_fill=BLUE):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    for i, width in enumerate(widths_cm):
        table.columns[i].width = Cm(width)
    header = table.rows[0]
    set_repeat_table_header(header)
    prevent_row_split(header)
    for i, text in enumerate(headers):
        set_cell_shading(header.cells[i], header_fill)
        set_cell_text(header.cells[i], text, bold=True, color=WHITE, size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER)
    for row in rows:
        cells = table.add_row().cells
        prevent_row_split(table.rows[-1])
        for i, text in enumerate(row):
            align = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT
            set_cell_text(cells[i], text, size=10, align=align)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_page_break(doc):
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


def build():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21)
    sec.page_height = Cm(29.7)
    sec.top_margin = Inches(1.0)
    sec.bottom_margin = Inches(1.0)
    sec.left_margin = Inches(1.25)
    sec.right_margin = Inches(1.25)
    sec.header_distance = Cm(1.2)
    sec.footer_distance = Cm(1.2)
    configure_styles(doc)

    # Cover
    for _ in range(2):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("中国研究生电子设计竞赛")
    set_run_font(r, "黑体", 20, bold=True)
    p.paragraph_format.space_after = Pt(28)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("技术论文撰写分工与项目完善建议")
    set_run_font(r, "黑体", 26, bold=True, color=BLUE)
    p.paragraph_format.space_after = Pt(18)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("面向煤矿安全规程审查与现场风险研判的\n多智能体系统")
    set_run_font(r, "黑体", 18, bold=True)
    p.paragraph_format.space_after = Pt(8)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("建议副标题：云边协同 · 工具增强 · 人机闭环")
    set_run_font(r, "楷体", 13, color="666666")

    for _ in range(5):
        doc.add_paragraph()
    meta = add_table(
        doc,
        ["项目", "建议填写内容"],
        [
            ["参赛单位", "____________________________"],
            ["团队成员", "成员A（统稿/智能体）  成员B（RAG/实验）  成员C（端侧/系统）"],
            ["指导教师", "____________________________"],
            ["编制日期", "2026年6月"],
        ],
        [4.0, 11.8],
        header_fill="5B9BD5",
    )
    for row in meta.rows[1:]:
        set_cell_shading(row.cells[0], LIGHT_GRAY)

    add_page_break(doc)

    # Executive abstract
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("摘  要")
    set_run_font(r, "黑体", 16, bold=True)
    add_body(
        doc,
        "本安排依据当前仓库最新 v9 实现和研电赛技术论文常用结构制定。项目已形成“结构保持型文档切分—本地混合检索—多链并行审查—确定性数值核验—争议升级—主智能体裁决—人工反馈—Word 改写—标注飞轮”的闭环。论文撰写不按平均字数切分，而按技术所有权、实验责任和证据材料划分；每位成员既负责章节，也负责对应代码、图表和演示验收。"
    )
    add_body(
        doc,
        "比赛前建议将微信小程序定位为移动交互入口，并增加真实嵌入式传感终端，组成“嵌入式采集节点—微信小程序—企业本地服务器”的云边协同原型。仅有微信小程序不能充分支撑“嵌入式平台”表述。"
    )
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0)
    r = p.add_run("关键词：")
    set_run_font(r, "黑体", 12, bold=True)
    r = p.add_run("煤矿安全；多智能体；检索增强生成；工具调用；微信小程序；嵌入式平台；人机闭环")
    set_run_font(r, "宋体", 12)

    add_callout(
        doc,
        "论文主线建议",
        [
            "作业前：审查作业规程，定位不合规内容并生成可追溯修改建议。",
            "作业中：现场人员通过小程序提交语音、图片和传感数据，智能体追问并调用工具研判风险。",
            "作业后：人工裁决与现场处置结果沉淀为高质量标注，持续改进检索和审查策略。",
        ],
        fill=PALE_GREEN,
    )

    # Chapter 1
    doc.add_heading("第一章 当前项目核验与参赛定位", level=1)
    doc.add_heading("1.1 当前仓库已实现能力", level=2)
    confirmed_rows = [
        ["结构化切分", "规则文档与待审文档分版本切分，保留章节、父级上下文和来源映射。", "chapter_based_chunking_v6.py；pending_doc_chunking_v9.py"],
        ["本地检索", "BGE-M3 Dense + Learned Sparse，RRF 融合并由 bge-reranker-v2-m3 重排；内容哈希缓存。", "hybrid_rag_review_v9.py"],
        ["并行审查", "合规链、错别字链、文档级重复链并行；chunk 级并发。", "hybrid_rag_review_v9.py"],
        ["工具增强", "LLM 抽取数值关系，代码执行单位归一化、阈值方向比较和反向漏报核验。", "numeric_compare.py"],
        ["自主升级", "分歧、数值冲突、低置信、异常和反向核验项进入 SQLite 升级队列。", "review_queue.py；main_agent_v9.py"],
        ["主智能体", "可调用审查、检索、数值比较、队列裁决等工具，并将裁决沉淀为标注。", "main_agent_v9.py"],
        ["Web 闭环", "SSE 实时推送问题；人工接受/驳回/改写；主智能体改写 Word 工作副本。", "v9_worker.py；v9_review.py；Review/index.jsx"],
        ["数据飞轮", "人工反馈作为最高优先级标注，可导出扩充黄金集。", "review_queue.py；v9_review.py"],
    ]
    add_table(doc, ["能力", "已确认实现", "主要证据"], confirmed_rows, [2.5, 8.0, 5.3])

    doc.add_heading("1.2 参赛作品建议定位", level=2)
    add_body(
        doc,
        "建议将作品定位为“面向煤矿安全规程审查与现场风险研判的云边协同多智能体系统”。不要只描述为文档审核工具，也不要把文档审查与硬件做成两个无关演示。二者应通过同一安全知识库、同一任务编排智能体和同一反馈飞轮连接起来。"
    )
    add_callout(
        doc,
        "赛题要点与项目对应关系",
        [
            "任务理解：主智能体识别“文档审查、数值核验、现场问答、风险处置”等任务并选择流程。",
            "工具调用：调用本地检索、数值比较、文档定位改写、传感器读取和告警工具。",
            "自主决策：低风险自动答复，存在分歧或高风险时主动追问、升级人工并生成处置建议。",
            "电子信息系统：FastAPI、SQLite 队列、SSE、Word 适配器、小程序与设备通信共同构成工程系统。",
            "嵌入式平台：真实传感节点或边缘终端负责数据采集、离线缓存和告警；小程序仅作为移动交互入口。",
        ],
    )

    doc.add_heading("1.3 论文中必须如实表述的边界", level=2)
    boundary_rows = [
        ["可直接写入“已实现”", "v9 审查引擎、BGE 本地检索、数值核验工具、升级队列、主智能体、人工反馈、Word 改写、数据飞轮。"],
        ["完成后再写入“已实现”", "微信小程序、语音/图片多模态输入、传感器采集、现场风险对话、真实嵌入式终端。"],
        ["不得表述为“完全本地”", "当前检索模型本地运行，但审查与改写使用 DashScope qwen-plus；应写成“本地检索 + 企业/云端大模型”。"],
        ["不得回避的限制", "worker 当前单文档串行；矿井类型索引缓存存在切换限制；复杂 Word 排版与连续编辑定位仍需增强。"],
    ]
    add_table(doc, ["边界类别", "具体内容"], boundary_rows, [4.1, 11.7], header_fill="C65911")

    # Chapter 2
    doc.add_heading("第二章 三人论文撰写总分工", level=1)
    doc.add_heading("2.1 分工原则", level=2)
    add_compact_item(doc, "原则一：", "一人统稿但不包办。成员A掌握全文论证与术语一致性，成员B、C对自己章节的事实和实验负责。")
    add_compact_item(doc, "原则二：", "章节、图表、代码证据、实验数据绑定到同一负责人，避免出现“文字写完但没有图和数据”的情况。")
    add_compact_item(doc, "原则三：", "每个核心创新至少有一个对照实验或系统验收；所有截图、指标和结论必须能够从当前代码或完成后的新模块复现。")

    doc.add_heading("2.2 三人角色与总交付物", level=2)
    roles = [
        ["成员A\n统稿/智能体", "项目总架构、任务理解、工具调用、自主决策、人机协同闭环", "摘要；第1、2章；第3章总架构；第5章智能体流程；第7章；成果价值", "系统总架构图、主智能体状态机图、端到端闭环图、创新点对照表", "全文术语/编号/引用一致；整合终稿与答辩叙事"],
        ["成员B\nRAG/算法实验", "结构化切分、本地混合检索、重排、数值核验、黄金集与评价", "第3章检索与工具；第4章算法原理；第6章算法实验", "切分示意图、检索漏斗图、三版指标表、数值核验对照表、消融实验", "所有指标脚本、数据口径、复现实验记录"],
        ["成员C\n端侧/系统工程", "FastAPI/worker/前端、Word 定位改写、微信小程序、嵌入式终端、系统测试", "第3章端云方案；第4章硬件；第5章软件实现；第6章系统测试", "部署图、硬件框图、电路/接线图、小程序流程图、页面截图、时延表", "完成移动端与嵌入式演示闭环；整理演示视频素材"],
    ]
    add_table(doc, ["角色", "技术主责", "论文主责", "必须提交的图表", "最终验收"], roles, [2.6, 3.5, 3.4, 3.7, 2.6])

    # Chapter 3
    doc.add_heading("第三章 按参考论文结构划分章节", level=1)
    add_body(
        doc,
        "下表沿用参考论文的“研究背景—难点创新—方案设计—原理硬件—软件流程—系统测试—总结—成果价值”组织方式。主笔负责首稿，协作人负责提供本专业证据，成员A负责最终统稿。"
    )
    chapter_rows = [
        ["摘要、关键词", "成员A", "成员B、C", "一句话场景、三项核心创新、系统闭环、关键实验结果；最后撰写。"],
        ["第一章 研究背景和意义", "成员A", "成员C", "煤矿规程审查痛点、现场安全问答痛点、现有方案不足、作品价值。"],
        ["第二章 作品难点与创新", "成员A", "成员B、C", "长文档跨条款审查、数值方向判断、低置信升级、端侧弱网；对应四项创新。"],
        ["第三章 方案论证与设计", "成员A", "成员B、C", "总体架构与任务闭环；B写 RAG/工具方案；C写小程序和嵌入式方案。"],
        ["第四章 原理分析及硬件电路图", "成员C", "成员B", "嵌入式采集节点、电源/传感器/通信；B补充 BGE、RRF、重排和数值核验原理。"],
        ["第五章 软件设计与流程", "成员C", "成员A、B", "Web/worker/SQLite/SSE/Word 改写；小程序；主智能体 ReAct/升级流程；数据飞轮。"],
        ["第六章 系统测试与分析", "成员B", "成员A、C", "B负责算法指标；A负责智能体与人工闭环；C负责端侧通信、时延、稳定性和端到端演示。"],
        ["第七章 总结", "成员A", "成员B、C", "总结已实现能力、实验结论、限制和下一步。"],
        ["研究成果及应用价值", "成员A", "全员", "工程成果、可迁移价值、应用边界、知识产权/软著/论文/竞赛材料。"],
        ["参考文献与附录", "成员B", "成员A、C", "B统一文献格式；附录放接口、硬件清单、实验配置和复现说明。"],
    ]
    add_table(doc, ["章节", "主笔", "协作", "写作与证据要求"], chapter_rows, [3.6, 1.8, 2.0, 8.8])

    doc.add_heading("3.1 成员A详细任务：统稿与主智能体", level=2)
    a_tasks = [
        ["A-1", "确定题目、摘要、关键词和统一术语表", "题目不夸大“全本地/嵌入式/多模态”；全文统一使用 v9 真实能力。"],
        ["A-2", "绘制系统总架构与业务闭环", "展示作业前审查、作业中研判、作业后飞轮，以及主智能体如何路由任务。"],
        ["A-3", "写任务理解、工具调用、自主决策机制", "基于 main_agent_v9.py 和升级队列，补齐现场对话决策流程后再写。"],
        ["A-4", "设计智能体实验", "对比无工具、仅 LLM、LLM+数值工具、LLM+升级/人工四种设置。"],
        ["A-5", "统稿与答辩主线", "统一图号、表号、引用、创新点和结论；准备 3 分钟系统故事线。"],
    ]
    add_table(doc, ["编号", "任务", "验收标准"], a_tasks, [1.7, 5.2, 9.0], header_fill="4472C4")

    doc.add_heading("3.2 成员B详细任务：RAG、工具与实验", level=2)
    b_tasks = [
        ["B-1", "写结构保持型切分与检索方案", "说明章节上下文、父子结构、父块保底+片段查询补充、矿井适用性过滤。"],
        ["B-2", "写 BGE-M3 + RRF + reranker 原理", "给出检索流程、关键参数、缓存机制和本地部署意义。"],
        ["B-3", "写“LLM抽取、代码比较”数值工具", "用方向语义、单位归一化、反向核验案例解释为什么优于纯 LLM。"],
        ["B-4", "建立统一黄金集与实验口径", "不能继续用各版本自身候选池直接做排行榜；统一证据组后再比较。"],
        ["B-5", "完成算法实验和误差分析", "至少报告 Hit/Recall/NDCG、问题级 P/R/F1、数值判断准确率、漏报/误报案例。"],
    ]
    add_table(doc, ["编号", "任务", "验收标准"], b_tasks, [1.7, 5.2, 9.0], header_fill="70AD47")

    doc.add_heading("3.3 成员C详细任务：系统、移动端与嵌入式", level=2)
    c_tasks = [
        ["C-1", "完善 v9 Web 闭环", "上传、SSE 实时问题、定位高亮、接受/驳回/改写、下载文档、飞轮查看稳定可演示。"],
        ["C-2", "开发微信小程序", "支持任务创建、语音/文字/图片上报、追问对话、风险卡片、人工确认和历史记录。"],
        ["C-3", "增加真实嵌入式采集节点", "建议 ESP32-S3 + 温湿度/气体模拟量/蜂鸣器或同类器件；小程序通过 BLE/Wi-Fi 获取数据。"],
        ["C-4", "实现现场风险工具接口", "主智能体可调用 get_sensor_reading、get_rule_evidence、raise_alarm、create_record 等工具。"],
        ["C-5", "完成工程测试", "报告端到端时延、弱网/断网缓存、重连、连续任务、编辑成功率和资源占用。"],
    ]
    add_table(doc, ["编号", "任务", "验收标准"], c_tasks, [1.7, 5.2, 9.0], header_fill="ED7D31")

    # Chapter 4
    doc.add_heading("第四章 必须准备的图、表、实验与材料", level=1)
    doc.add_heading("4.1 建议图件清单", level=2)
    figs = [
        ["图1", "系统总体架构图", "成员A", "嵌入式节点—小程序—API/队列—主智能体—审查工具—知识库—反馈飞轮。"],
        ["图2", "端到端业务闭环图", "成员A", "作业前、作业中、作业后形成连续安全闭环。"],
        ["图3", "结构保持型切分示意图", "成员B", "标题层级、父上下文、规则块与待审块映射。"],
        ["图4", "混合检索与多查询流程图", "成员B", "父块查询、片段查询、Dense/Sparse、RRF、重排、去重。"],
        ["图5", "数值核验工具流程图", "成员B", "LLM 抽取—代码归一化—方向比较—反向漏报核验—升级。"],
        ["图6", "主智能体与升级队列状态图", "成员A", "自动裁决、主动追问、人工升级、dead-letter。"],
        ["图7", "嵌入式终端硬件框图与接线图", "成员C", "传感器、MCU、电源、通信、告警。"],
        ["图8", "微信小程序页面与交互流程", "成员C", "现场上报、追问、风险卡、确认、历史记录。"],
        ["图9", "Web 审查与 Word 改写截图", "成员C", "问题定位、人工反馈、改写后文档与下载。"],
        ["图10", "数据飞轮与持续评估流程", "成员A/B", "人工标注导出、黄金集、指标回归、版本更新。"],
    ]
    add_table(doc, ["编号", "图件", "负责人", "必须表达的信息"], figs, [1.3, 3.8, 2.0, 8.8])

    doc.add_heading("4.2 建议实验矩阵", level=2)
    experiments = [
        ["E1 检索对比", "旧混合检索 / 父块 BGE / 父块+片段双通道", "Hit@K、Recall@K、MRR、NDCG、覆盖率", "成员B"],
        ["E2 数值核验", "纯 LLM / LLM+正向工具 / LLM+正反向工具", "准确率、漏报率、方向误判率", "成员B"],
        ["E3 智能体消融", "无升级 / 自动升级 / 自动升级+人工", "最终问题 P/R/F1、人工介入率、平均步骤数", "成员A/B"],
        ["E4 Word 改写", "段落、表格、多处同文、连续编辑、复杂文档", "定位成功率、改写成功率、格式保持率", "成员C"],
        ["E5 端侧链路", "在线、弱网、断网缓存、重连", "上报时延、成功率、恢复时间、设备资源占用", "成员C"],
        ["E6 完整演示", "规程审查→现场上报→智能体追问→工具研判→人工确认→飞轮", "端到端时延、流程完成率、可追溯记录", "全员"],
    ]
    add_table(doc, ["实验", "对照设置", "核心指标", "负责人"], experiments, [2.5, 5.3, 5.8, 2.2], header_fill="5B9BD5")

    doc.add_heading("4.3 当前可用的实验事实与使用限制", level=2)
    add_body(
        doc,
        "现有三版检索对比记录显示，当前 BGE-M3 证据组版在 50 个案例中有 43 个案例出现 useful 证据，覆盖率为 86%；但其 MRR、Hit@1 和 NDCG@15 并不占优。该结果可以用于说明“覆盖率提升与前排排序仍需同时优化”，不能直接宣称 BGE 版全面优于旧版。"
    )
    add_body(
        doc,
        "当前不同检索版本使用各自候选池中的人工标签，黄金集并非严格同源。正式论文应先建立独立统一证据组黄金集，再给出模型排行榜；否则需要在表注中明确其仅用于候选覆盖和排序诊断。"
    )

    # Chapter 5
    doc.add_heading("第五章 下一步完善项目的优先级建议", level=1)
    doc.add_heading("5.1 P0：比赛主线必须补齐", level=2)
    p0 = [
        ["P0-1", "小程序 + 真实嵌入式节点", "把小程序作为人机交互入口，把 ESP32-S3/同类设备作为采集与告警平台。完成“设备数据进入智能体决策”的真实闭环。", "成员C", "可现场演示"],
        ["P0-2", "现场风险对话智能体", "增加任务分类、主动追问、传感器读取、法规检索、风险分级、告警/升级工具。复用现有主智能体与队列设计。", "成员A/C", "满足赛题任务理解/工具调用/自主决策"],
        ["P0-3", "统一黄金集与端到端评测", "固定同一批待审内容和证据组，补齐问题级审查指标、数值工具指标和系统指标。", "成员B", "论文核心数据可信"],
        ["P0-4", "形成完整演示脚本", "准备一条清晰场景：规程中阈值错误→审查发现→人工确认改写→现场设备上报异常→智能体追问并告警→记录进入飞轮。", "全员", "答辩故事完整"],
    ]
    add_table(doc, ["编号", "任务", "实施要点", "负责人", "完成标志"], p0, [1.5, 3.3, 7.0, 2.0, 2.2], header_fill="C00000")

    doc.add_heading("5.2 P1：工程质量与可信度", level=2)
    p1 = [
        ["P1-1", "按矿井类型维护独立索引", "当前 worker 启动时按 non_outburst 建索引并缓存，任务中切换 mine_type 不会真正重建专用索引。"],
        ["P1-2", "多任务并发与资源隔离", "当前 worker 单文档串行；增加多 worker、GPU 检索批处理和任务级限流。"],
        ["P1-3", "稳定 Word 定位协议", "连续编辑后 block index 可能漂移；建议引入稳定段落 ID/书签、原文哈希和版本号，复杂表格单独处理。"],
        ["P1-4", "可追溯审查证据", "每条问题保存法规来源、检索分数、工具结论、智能体步骤、人工裁决和最终改写。"],
        ["P1-5", "安全与权限", "增加用户/任务归属、文件类型与大小校验、审计日志、密钥隔离和接口访问控制。"],
        ["P1-6", "前端性能", "当前构建产物主 JS 约 1.9 MB；使用路由懒加载和代码分包，提升移动端加载体验。"],
    ]
    add_table(doc, ["编号", "任务", "原因与做法"], p1, [1.7, 4.1, 10.0], header_fill="BF9000")

    doc.add_heading("5.3 P2：加分项与后续扩展", level=2)
    add_compact_item(doc, "P2-1：", "增加语音转写、现场照片识别和多轮追问，形成真正的多模态现场助手。")
    add_compact_item(doc, "P2-2：", "把高频风险规则和轻量分类模型下沉到边缘侧，断网时仍能做基础阈值告警与任务缓存。")
    add_compact_item(doc, "P2-3：", "增加规则版本管理、法规更新影响分析和历史规程回溯，强化工程应用价值。")
    add_compact_item(doc, "P2-4：", "围绕结构保持型切分、数值工具核验、升级飞轮申请软著/专利或形成论文成果。")

    doc.add_heading("5.4 最小可行嵌入式方案", level=2)
    embedded_rows = [
        ["嵌入式采集节点", "ESP32-S3 或同类 MCU 开发板；温湿度传感器；可燃气体/甲烷模拟量模块（比赛原型）；蜂鸣器/LED；电池供电。", "体现真实嵌入式数据采集、边缘阈值告警和通信。"],
        ["移动交互入口", "微信小程序：语音/文字/图片上报，BLE/Wi-Fi 读取设备数据，显示风险卡片并确认处置。", "适合展示，不应单独称为嵌入式平台。"],
        ["智能服务端", "现有 FastAPI + SQLite 队列 + v9 worker + 主智能体 + 本地 BGE 检索。", "最大程度复用当前成果。"],
        ["演示通信", "比赛现场使用 BLE 或局域网；论文中将矿井生产部署表述为需接入矿用通信网络和防爆/本安设备。", "避免把竞赛原型误写成可直接下井产品。"],
    ]
    add_table(doc, ["组成", "建议实现", "定位"], embedded_rows, [3.2, 8.7, 3.9], header_fill="548235")

    # Chapter 6
    doc.add_heading("第六章 协作节奏与验收机制", level=1)
    doc.add_heading("6.1 建议十日冲刺节奏", level=2)
    schedule = [
        ["D1", "冻结题目、主线、术语和论文目录；确定嵌入式器件。", "A主责，全员确认"],
        ["D2-D3", "C完成小程序骨架和设备通信；B冻结统一实验集；A完成架构图与第一、二章。", "并行推进"],
        ["D4-D5", "接通现场风险工具；完成主要算法实验；补齐第三至第五章首稿。", "各自技术主责"],
        ["D6", "完整演示联调；记录失败案例、时延、截图和视频。", "全员"],
        ["D7-D8", "完成第六章实验分析；补齐硬件图、流程图、表格和引用。", "B/C主责，A统稿"],
        ["D9", "交叉审稿：事实、指标、格式、图表编号、创新点逐项核验。", "A审B、B审C、C审A"],
        ["D10", "终稿、答辩 PPT、演示视频与故障兜底方案。", "全员"],
    ]
    add_table(doc, ["时间", "主要任务", "责任"], schedule, [2.0, 10.6, 3.2])

    doc.add_heading("6.2 每日交付与交叉审稿", level=2)
    add_compact_item(doc, "每日提交：", "新增代码、可复现命令、实验结果 JSON/CSV、论文新增段落、图表源文件、问题清单。")
    add_compact_item(doc, "成员A检查：", "是否回答赛题、是否体现任务理解/工具调用/自主决策、是否存在夸大表述。")
    add_compact_item(doc, "成员B检查：", "指标口径、对照公平性、实验可复现性、结论是否由数据支持。")
    add_compact_item(doc, "成员C检查：", "系统链路能否真实运行、截图与当前界面一致、嵌入式与移动端是否能现场演示。")

    doc.add_heading("6.3 最终验收清单", level=2)
    checks = [
        ["论文", "参考论文式章节完整；图表编号连续；每项创新均有原理、实现、实验和限制。"],
        ["系统", "Web 审查闭环稳定；小程序与嵌入式数据真实接通；完整场景可连续演示。"],
        ["实验", "统一黄金集；算法、智能体、Word 改写、端侧链路和端到端测试均有记录。"],
        ["材料", "架构图、流程图、硬件图、页面截图、演示视频、复现说明、代码版本号齐全。"],
        ["表述", "不把小程序单独称为嵌入式；不把 DashScope 推理称为完全本地；明确原型与生产部署边界。"],
    ]
    add_table(doc, ["类别", "验收标准"], checks, [3.0, 12.8], header_fill="4472C4")

    # Chapter 7
    doc.add_heading("第七章 结论", level=1)
    add_body(
        doc,
        "当前 v9 已具备较强的参赛基础，真正有辨识度的创新不是简单叠加多个智能体，而是将本地混合检索、确定性数值工具、争议升级、人工裁决、Word 改写和标注飞轮组织成可信闭环。下一步应优先补齐现场风险对话与真实嵌入式节点，并用统一黄金集和端到端实验把系统能力量化。"
    )
    add_body(
        doc,
        "三人分工必须以“技术模块—论文章节—图表实验—演示验收”四项绑定为原则。成员A保证论证和智能体主线，成员B保证算法与指标可信，成员C保证移动端、嵌入式和工程闭环可运行。这样既能降低重复劳动，也能让答辩时每位成员都能对自己负责的技术给出可验证证据。"
    )

    # Footer
    for section in doc.sections:
        footer = section.footer
        p = footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run("研电赛论文撰写分工与项目完善建议  |  ")
        set_run_font(r, "宋体", 9, color="777777")
        add_field(p, "PAGE")

    doc.core_properties.title = "研电赛论文三人分工与项目完善建议"
    doc.core_properties.subject = "煤矿安全规程审查与现场风险研判多智能体系统"
    doc.core_properties.author = "项目团队"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
