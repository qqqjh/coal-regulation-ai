from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement


DOCX = Path("D:/work/project/coal-regulation-ai/" + "\u5206\u5de5_\u4e09\u4eba\u6574\u5408\u7248_xml.docx")


def set_paragraph_text(paragraph, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def insert_paragraph_after(paragraph, text: str):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = paragraph._parent.add_paragraph()
    new_para._p = new_p
    new_para.style = paragraph.style
    new_para.add_run(text)
    return new_para


def main() -> None:
    doc = Document(str(DOCX))

    old = (
        "\u4efb\u52a1\u6570\u636e\u5c42\u4ee5 SQLite \u961f\u5217\u4e3a\u6838\u5fc3"
    )
    replacement = (
        "\u4efb\u52a1\u6570\u636e\u5c42\u4ee5 MongoDB \u4efb\u52a1\u96c6\u5408\u4e0e Redis "
        "\u7f13\u5b58\u961f\u5217\u4e3a\u6838\u5fc3\uff0c\u8d1f\u8d23\u8fde\u63a5\u4e1a\u52a1"
        "\u670d\u52a1\u4e0e\u5e38\u9a7b worker\u3002\u4e1a\u52a1\u670d\u52a1\u5c42\u5728"
        "\u7528\u6237\u4e0a\u4f20\u6587\u6863\u540e\u521b\u5efa\u5f85\u5904\u7406\u4efb\u52a1"
        "\uff0cworker \u901a\u8fc7\u72b6\u6001\u8f6e\u8be2\u6216\u961f\u5217\u6d88\u8d39"
        "\u9886\u53d6\u4efb\u52a1\u5e76\u56de\u5199\u72b6\u6001\u3002MongoDB \u4fdd\u5b58"
        "\u4efb\u52a1\u8fdb\u5ea6\u3001\u6bb5\u843d\u6a21\u578b\u3001\u95ee\u9898\u8bb0"
        "\u5f55\u3001\u4eba\u5de5\u53cd\u9988\u3001\u5347\u7ea7\u9879\u548c\u7ecf\u9a8c"
        "\u6807\u6ce8\uff1bRedis \u7528\u4e8e\u7f13\u5b58\u4efb\u52a1\u72b6\u6001"
        "\u3001\u589e\u91cf\u6d88\u606f\u548c\u77ed\u671f\u4f1a\u8bdd\uff0c\u652f\u6491 "
        "SSE \u63a8\u9001\u3001\u5931\u8d25\u91cd\u8bd5\u548c\u9ad8\u9891\u67e5\u8be2"
        "\u3002\u901a\u8fc7\u201cMongoDB \u6301\u4e45\u5316 + Redis \u77ed\u671f"
        "\u7f13\u5b58\u201d\u7684\u65b9\u5f0f\uff0c\u7cfb\u7edf\u80fd\u591f\u652f"
        "\u6301\u4e1a\u52a1\u670d\u52a1\u548c worker \u8de8\u8fdb\u7a0b\u534f\u4f5c"
        "\uff0c\u5b9e\u73b0\u4efb\u52a1\u65ad\u70b9\u6062\u590d\u3001\u95ee\u9898"
        "\u589e\u91cf\u6d41\u548c\u53cd\u9988\u5904\u7406\u961f\u5217\u3002"
    )
    for para in doc.paragraphs:
        if para.text.startswith(old):
            set_paragraph_text(para, replacement)
            break

    marker = (
        "\u6570\u636e\u5b58\u50a8\u5c42\u7531 MongoDB \u4e0e Milvus \u7ec4\u6210"
    )
    insert_text = (
        "\u77e5\u8bc6\u5904\u7406\u4e0e\u6a21\u578b\u670d\u52a1\u5c42\u4e2d\uff0c"
        "MinerU \u7528\u4e8e\u5bf9\u626b\u63cf\u7248 PDF\u3001\u56fe\u7247\u578b"
        "\u6587\u6863\u548c\u590d\u6742\u7248\u5f0f\u6750\u6599\u8fdb\u884c OCR "
        "\u4e0e\u7248\u9762\u89e3\u6790\uff0cBGE-M3 \u540c\u65f6\u627f\u62c5"
        "\u6587\u672c\u5d4c\u5165\u4e0e\u91cd\u6392\u8bc4\u5206\uff0cQwen \u5927"
        "\u6a21\u578b\u627f\u62c5\u4e3b\u667a\u80fd\u4f53\u89c4\u5212\u3001\u4e13"
        "\u4e1a\u667a\u80fd\u4f53\u5ba1\u67e5\u3001\u8ffd\u95ee\u751f\u6210\u548c"
        "\u7ed3\u679c\u5f52\u7eb3\u3002Redis \u4f5c\u4e3a\u7f13\u5b58\u4e0e"
        "\u5f02\u6b65\u6d88\u606f\u7f13\u51b2\u7ec4\u4ef6\uff0c\u7528\u4e8e\u4fdd"
        "\u5b58\u70ed\u70b9\u68c0\u7d22\u7ed3\u679c\u3001\u4efb\u52a1\u72b6\u6001"
        "\u3001SSE \u589e\u91cf\u8f93\u51fa\u548c\u77ed\u671f\u4f1a\u8bdd\u4e0a"
        "\u4e0b\u6587\uff0c\u964d\u4f4e\u91cd\u590d\u68c0\u7d22\u4e0e\u6a21\u578b"
        "\u8c03\u7528\u5f00\u9500\u3002"
    )
    for para in doc.paragraphs:
        if para.text.startswith(marker):
            insert_paragraph_after(para, insert_text)
            break

    doc.save(str(DOCX))


if __name__ == "__main__":
    main()
