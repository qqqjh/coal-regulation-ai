from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path(r"D:\work\项目\研电赛\分工_第二章格式修改版.docx")
OUTPUT = ROOT / "分工_第三章格式扩写版.docx"


INTRO = (
    "本项目围绕煤矿作业规程智能审查过程中“规则知识难统一、长篇文档难检索、复杂任务难协同、"
    "有效经验难复用”四类核心问题，构建从知识组织、证据检索、智能决策到持续进化的完整技术方案。"
    "首先，面向法规、规程、标准和实施细则等异构规则文档，建立兼顾章节层级、条款语义、适用范围和"
    "版本来源的多结构规则知识库；其次，针对待审文档篇幅长、主题多和关键要求易被稀释的问题，设计"
    "结构感知切分与多查询协同检索方法，为审查智能体提供可定位、可追溯的法规证据。在决策层，采用"
    "主智能体统筹、专业智能体异步并行、确定性工具辅助核验和人工兜底的协同机制，提高审查效率与"
    "结论可信度。最后，通过提示词、Skill、案例记忆和人工反馈数据飞轮分层沉淀经验，使系统在安全"
    "可控的前提下持续优化，并为后续跨文档、跨矿井应用提供可扩展基础。"
)


SECTIONS = [
    (
        "3.1 多结构规则知识库构建（成员B，画图）",
        [
            "煤矿安全审查所依据的知识来源包括国家法规、行业标准、企业规程和专项实施细则。不同文档在篇章结构、编号方式、条款粒度和适用对象方面差异明显，若采用固定字符长度切分，容易割裂条款之间的上下位关系，并导致限制条件、例外情形和适用范围丢失。为此，本项目设计面向异构规则文档的多结构知识库构建方案，将文档内容表示为“文档结构树、可检索证据单元和规则元数据”三类相互关联的信息。",
            "在知识加工阶段，系统首先通过版面解析与标题识别恢复文档的篇、章、节、条及编号项结构，再依据语义完整性和最大长度约束生成证据单元。每个证据单元均保留父级标题链、相邻条款、原始页码、来源文件、发布时间和版本号等信息。针对突出矿井、特定作业面和专用设备等适用性要求，进一步建立场景标签与规则约束字段，使检索过程能够在语义相关性之外执行适用范围过滤，避免将不适用条款作为审查依据。",
            "在知识表示与索引阶段，方案采用层级结构索引与混合向量索引协同的方式。结构索引用于维护章节关系、规则引用和版本溯源；混合向量索引利用 BGE-M3 同时生成语义稠密向量与学习型稀疏权重，以兼顾专业术语精确匹配和语义相似召回。对于同一规则在不同文档中的重复表述，系统按证据组进行关联与去重，并保留权威来源和最新版本。该设计能够支持知识库增量更新、法规版本切换和审查证据追溯，具备较好的工程可行性与扩展能力。",
            "为保证入库知识质量，系统设置结构完整性、来源可信度、适用性标签和可检索性四类质量检查。无法识别层级或存在版本冲突的条款进入人工确认队列，确认后再发布到正式知识库。通过上述机制，知识库由单纯的文本片段集合升级为具有结构、语义、场景和版本信息的规则证据体系，为后续检索与智能审查提供可靠基础。",
        ],
        "图3-1  多结构规则知识库构建总体流程",
        "建议绘图：异构规则文档→版面与章节识别→结构化证据单元→场景/版本标注→结构索引与混合向量索引→质量校验与发布",
    ),
    (
        "3.2 待审文档结构感知切分与多查询检索（成员B，画图）",
        [
            "待审作业规程通常包含大量章节、表格、编号项和跨段落约束，单个章节内部还可能同时描述作业条件、操作步骤、安全阈值和应急措施。若直接以整篇文档或超长段落作为查询，关键违规内容的语义信号容易被其他主题稀释；若切分过细，则可能失去作业场景和父级章节约束。针对上述矛盾，本项目提出结构感知的父子切分与多查询协同检索方案。",
            "切分阶段以标题层级和段落边界为基础建立父级语义块，并将其中的编号项、关键句和数值要求拆解为子级检索单元。父级块负责保存完整场景、章节上下文和原文位置，子级单元负责突出具体动作、对象、条件和阈值。对于包含多个安全要求的复杂段落，主智能体进一步将其分解为若干原子审查主张，并为每条主张生成包含场景词、动作词和参数词的检索查询，从而减少多主题内容对证据召回的干扰。",
            "检索阶段采用“父级整体语义保底、子级与原子主张精准补充”的双通道机制。各查询分别经过 BGE-M3 稠密召回、学习型稀疏召回和倒数排名融合，再使用交叉编码重排序模型评估待审主张与规则证据之间的相关性。系统同时利用矿井类型、章节场景和规则有效版本等元数据进行过滤，并依据多查询覆盖数量、最佳重排分数和证据权威等级进行综合排序。最终候选按证据组去重，既避免重复条款占据候选位置，又保留必要的上下文证据。",
            "为提高审查过程的可解释性，每条召回证据均记录其对应的待审原文位置、触发查询、召回通道、相关性得分和来源条款。当检索置信度不足或不同查询结果相互冲突时，系统自动扩大候选范围、改写查询或提交主智能体复核。该方案能够在保持长文档结构语义的同时提高关键法规证据的覆盖率，并为问题定位、人工复核和 Word 修订提供稳定映射。",
        ],
        "图3-2  待审文档结构感知切分与多查询检索流程",
        "建议绘图：待审文档→父级语义块/子级单元/原子主张→多查询生成→Dense/Sparse/RRF→重排序与元数据过滤→证据组聚合与定位",
    ),
    (
        "3.3 主智能体自主决策、工具核验与多智能体异步并行审查（成员B，画图）",
        [
            "为解决多类型审查任务路径复杂、串行处理效率低以及低置信结论难以可靠裁决的问题，本项目设计基于 LangChain 工具封装与 LangGraph 状态编排的主智能体决策框架。主智能体接收审查任务后，首先识别文档类型、审查目标和风险等级，生成包含执行步骤、专业智能体分工、工具调用和升级条件的结构化策略，再将任务分发给知识检索、合规审查、错别字检查、重复性检查、数值核验和结果修订等专业智能体。",
            "审查流程采用“检索预计算与专业智能体异步并行”机制。系统先批量完成待审块的混合检索与重排序，随后并行启动合规审查链、错别字检查链和文档级重复性检查链，并在各链内部根据任务复杂度继续执行块级并发。各专业智能体在独立上下文中运行，通过统一结构化结果协议返回问题类型、原文位置、法规依据、置信度和修改建议，主智能体负责汇总、去重和冲突识别，从而降低上下文相互污染并缩短总体审查时间。",
            "对于数值阈值、单位换算和方向语义等语言模型容易误判的问题，系统采用“语言模型抽取、确定性工具比较”的核验机制。语言模型负责识别参数对象、数值、单位、比较方向和适用场景，数值工具执行单位归一化、上下限方向编码和严格程度比较。针对初审判断合规但包含报警、断电、复电或停工阈值的内容，系统执行反向漏报检查；发现疑似超限或方向错误时，不直接强制判定，而是进入争议升级流程。",
            "主智能体依据证据充分性、专业智能体一致性、工具核验结果和风险等级进行自主决策。普通一致结果自动合并输出；初审与复核分歧、工具与模型冲突、低置信结论和执行异常进入升级队列，由主智能体重新检索、追加工具调用或转交人工确认。整个状态流设置最大执行步数、异常重试、人工中断和结果审计机制，使智能体自主性始终处于可解释、可追踪和可控制的边界内。",
        ],
        "图3-3  主智能体自主决策、工具核验与多智能体异步并行审查流程",
        "建议绘图：任务输入→主智能体策略生成→检索预计算→专业智能体异步并行→工具核验→结果合并/冲突检测→升级队列→人工确认与修订",
    ),
    (
        "3.4 多层经验保存、提炼与持续进化机制（成员B，画图）",
        [
            "煤矿规程审查具有较强的专业性和场景差异，同类问题在不同矿井、不同文档和不同版本法规中可能呈现不同表达。若系统仅保存最终审查结果，难以复用稳定规则和高价值人工经验；若未经筛选地自动学习，又可能引入错误偏差。为此，本项目提出面向可信审查的多层经验保存、提炼与持续进化机制，将经验按照稳定程度、可信来源和复用范围划分为规则记忆、流程记忆、案例记忆和人工权威记忆。",
            "规则记忆主要保存经过验证的领域约束、提示词模板和输出规范，用于约束智能体的基础判断边界；流程记忆以 Skill 形式封装规则文档入库、待审文档检索、并行审查、争议升级和 Word 修订等可重复流程，使主智能体能够按任务读取并调用；案例记忆记录主智能体对冲突项的裁决过程、调用工具和最终理由，用于相似问题检索和错误复盘；人工权威记忆保存专家的接受、驳回和自定义修改意见，并作为最高优先级经验进入数据飞轮。",
            "经验提炼过程设置来源分级、置信度评估、重复聚类和人工确认等质量门控。系统定期从新增案例中识别高频错误模式和稳定处理策略，形成候选提示词规则或候选 Skill；只有经过人工确认和回归测试后，候选经验才能升级到正式版本。对于法规更新或实践证明不再适用的经验，系统通过版本标识、适用范围和回滚机制完成安全退出，避免历史经验污染当前审查结果。",
            "在持续评估阶段，人工确认案例与高置信裁决可用于扩充统一黄金标注集，并驱动检索效果、专业智能体功能和主智能体决策质量的回归测试。新版本只有在关键指标不下降且高风险案例通过复核后才能发布。通过“运行反馈—经验提炼—人工确认—回归评测—版本发布”的闭环，系统能够持续适应企业规程编制与审查需求，同时保持法规依据的权威性和智能决策的安全可控性。",
        ],
        "图3-4  多层经验保存、提炼与持续进化机制",
        "建议绘图：提示词规则/Skill流程/智能体案例/人工反馈→质量门控与经验提炼→黄金集与回归评测→版本发布/回滚→持续运行反馈",
    ),
]


def set_cell_shading(cell, fill="F2F2F2"):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def add_body_before(doc, anchor, text):
    p = doc.add_paragraph(style="Normal")
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Cm(0.74)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(12)
    anchor.addprevious(p._p)


def add_heading_before(doc, anchor, text):
    p = doc.add_paragraph(text, style="Heading 2")
    anchor.addprevious(p._p)


def add_figure_placeholder_before(doc, anchor, note, caption):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Cm(14.5)
    cell = table.cell(0, 0)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_shading(cell)
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(14)
    r = p.add_run(note)
    r.font.name = "Times New Roman"
    r._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
    r.font.size = Pt(10.5)
    r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    anchor.addprevious(table._tbl)

    cp = doc.add_paragraph(style="Caption")
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.first_line_indent = Pt(0)
    cp.paragraph_format.space_before = Pt(4)
    cp.paragraph_format.space_after = Pt(6)
    cr = cp.add_run(caption)
    cr.font.name = "Times New Roman"
    cr._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
    cr.font.size = Pt(10.5)
    anchor.addprevious(cp._p)


def find_heading(doc, prefix):
    for p in doc.paragraphs:
        if p.text.strip().startswith(prefix):
            return p
    raise ValueError(f"未找到标题: {prefix}")


def remove_between(start_p, end_p):
    node = start_p._p.getnext()
    while node is not None and node is not end_p._p:
        nxt = node.getnext()
        node.getparent().remove(node)
        node = nxt


def build():
    doc = Document(SOURCE)
    chapter3 = find_heading(doc, "第三章")
    chapter4 = find_heading(doc, "第四章")
    chapter3.text = "第三章 方案论证与设计"

    remove_between(chapter3, chapter4)
    anchor = chapter4._p

    add_body_before(doc, anchor, INTRO)
    for heading, paragraphs, caption, note in SECTIONS:
        add_heading_before(doc, anchor, heading)
        for paragraph in paragraphs:
            add_body_before(doc, anchor, paragraph)
        add_figure_placeholder_before(doc, anchor, note, caption)

    doc.core_properties.title = "分工文档第三章格式扩写版"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
