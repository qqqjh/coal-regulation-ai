from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "研电赛技术论文简要初稿_含三人分工标注_v7.docx"
INK = "222222"
BLUE = "1F4E79"
GRAY = "F2F2F2"


def font(run, east="宋体", size=12, bold=False, color=INK, italic=False):
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def no_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    tr_pr.append(node)


def cell_text(cell, text, bold=False, white=False, center=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(text)
    font(r, "宋体", 10.5, bold=bold, color="FFFFFF" if white else INK)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def table(doc, headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, w in enumerate(widths):
        t.columns[i].width = Cm(w)
    repeat_header(t.rows[0])
    no_split(t.rows[0])
    for i, h in enumerate(headers):
        shade(t.rows[0].cells[i], BLUE)
        cell_text(t.rows[0].cells[i], h, bold=True, white=True, center=True)
    for row in rows:
        cells = t.add_row().cells
        no_split(t.rows[-1])
        for i, val in enumerate(row):
            cell_text(cells[i], str(val), center=i == 0)
    doc.add_paragraph()
    return t


def body(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Cm(0.74)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(0)
    font(p.add_run(text), "宋体", 12)
    return p


def item(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.74)
    p.paragraph_format.first_line_indent = Cm(-0.74)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(0)
    font(p.add_run(text), "宋体", 12)


def caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(6)
    font(p.add_run(text), "宋体", 10.5)


def placeholder(doc, text):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    t.columns[0].width = Cm(14.5)
    c = t.cell(0, 0)
    shade(c, GRAY)
    cell_text(c, "\n" + text + "\n", center=True)
    doc.add_paragraph()


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def setup(doc):
    sec = doc.sections[0]
    sec.page_width = Cm(21)
    sec.page_height = Cm(29.7)
    sec.top_margin = Inches(1)
    sec.bottom_margin = Inches(1)
    sec.left_margin = Inches(1.25)
    sec.right_margin = Inches(1.25)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE

    for name, east, size, align in [
        ("Heading 1", "黑体", 18, WD_ALIGN_PARAGRAPH.CENTER),
        ("Heading 2", "黑体", 15, WD_ALIGN_PARAGRAPH.LEFT),
        ("Heading 3", "黑体", 12, WD_ALIGN_PARAGRAPH.LEFT),
    ]:
        s = doc.styles[name]
        s.font.name = "Times New Roman"
        s._element.rPr.rFonts.set(qn("w:eastAsia"), east)
        s.font.size = Pt(size)
        s.font.bold = True
        s.font.color.rgb = RGBColor.from_string(INK)
        s.paragraph_format.alignment = align
        s.paragraph_format.first_line_indent = Cm(0)
        s.paragraph_format.space_before = Pt(10)
        s.paragraph_format.space_after = Pt(5)
        s.paragraph_format.line_spacing = 1.5


def build():
    doc = Document()
    setup(doc)

    # Cover
    for _ in range(2):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font(p.add_run("第十九届中国研究生电子设计竞赛技术论文"), "黑体", 20, bold=True)
    p.paragraph_format.space_after = Pt(38)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font(p.add_run("面向煤矿安全规程审查与现场风险研判的\n移动端协同多智能体系统"), "黑体", 24, bold=True, color=BLUE)
    p.paragraph_format.space_after = Pt(10)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font(p.add_run("A Cloud-Edge Collaborative Multi-Agent System for Coal Mine Safety"), "Times New Roman", 13, italic=True)
    for _ in range(5):
        doc.add_paragraph()
    table(doc, ["项  目", "内  容"], [
        ["参赛单位", "____________________________"],
        ["团队成员", "成员A、成员B、成员C（请替换为真实姓名）"],
        ["指导教师", "____________________________"],
        ["作品类别", "面向大语言模型、多模态模型及智能体系统的创新设计"],
    ], [4.0, 11.8])

    page_break(doc)

    # Abstract
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font(p.add_run("摘  要（成员A主笔，成员B/C提供数据）"), "黑体", 16, bold=True)
    body(doc, "煤矿作业规程篇幅长、条款关联复杂，传统人工审查存在效率低、数值阈值易误判、问题定位与修订闭环不足等问题。同时，井下现场风险信息具有突发性和多源性，需要能够理解任务、主动追问并调用工具的智能系统。为此，本作品设计了一套面向煤矿安全规程审查与现场风险研判的移动端协同多智能体系统。")
    body(doc, "系统在规程审查侧采用结构保持型文档切分、本地 BGE-M3 稠密与稀疏混合检索、重排序和多查询召回，通过合规审查、错别字检查和文档级重复检查等并行链路完成审查；针对数值阈值问题，采用“语言模型抽取、确定性代码比较”的工具增强方法，并将争议项自动升级至主智能体或人工复核。系统支持问题实时定位、人工接受或改写、Word 文档修订及标注数据回流。")
    body(doc, "在现场应用侧，系统通过微信小程序调用智能手机已有的麦克风、摄像头、定位、存储和无线通信能力，实现语音、文字与图片上报。主智能体根据现场任务调用法规检索、风险分析和记录保存等工具，在信息不足时主动追问，在高风险或低置信情况下升级人工。该系统形成“作业前审查、作业中研判、作业后沉淀”的安全闭环。")
    p = doc.add_paragraph()
    font(p.add_run("关键词："), "黑体", 12, bold=True)
    font(p.add_run("煤矿安全；多智能体；检索增强生成；工具调用；微信小程序；移动智能终端"), "宋体", 12)

    page_break(doc)

    # Chapter 1
    doc.add_heading("第一章 研究背景和意义（成员A主笔，成员C协作）", level=1)
    body(doc, "煤矿作业规程是指导生产组织和安全作业的重要技术文件，其内容涉及通风、瓦斯、防治水、顶板管理、机电运输等多个专业。随着法规标准不断更新，规程编制和审查需要同时考虑章节结构、适用场景、数值阈值和跨条款依赖，单纯依赖人工逐条核对容易产生漏审和误判。")
    body(doc, "大语言模型具备较强的文本理解和生成能力，但直接使用模型进行合规判断仍存在法规依据召回不充分、数值方向判断不稳定和结果不可追溯等问题。因此，需要将检索增强生成、确定性工具、人机协同和持续评估机制结合起来，提高审查结果的可信度。")
    body(doc, "此外，煤矿现场风险往往需要结合人员描述和现场图像综合判断。将规程审查形成的安全知识用于微信小程序中的现场问答与风险研判，可形成从规程到执行、从执行到反馈的闭环。")
    placeholder(doc, "待补图：煤矿规程人工审查与现场风险处置的主要痛点")
    caption(doc, "图1-1  项目应用背景与需求")

    # Chapter 2
    doc.add_heading("第二章 作品难点与创新（成员A统稿，成员B/C分项撰写）", level=1)
    doc.add_heading("2.1 异构规则文档的结构统一与知识组织难题（成员B主笔）", level=2)
    body(doc, "煤矿法规、规程、国家标准和实施细则在篇章层级、编号形式、内容粒度和适用范围方面存在明显差异。若仅按固定字符长度切分，容易破坏条款结构、丢失父级语境，并将突出矿井专用要求错误用于一般场景。因此，如何保留规则文档结构并形成统一、可检索的多结构知识库，是系统首先需要解决的难题。")
    doc.add_heading("2.2 长篇待审文档的语义保持与精准证据检索难题（成员B主笔）", level=2)
    body(doc, "待审作业规程篇幅长、章节层级复杂，单个内容块中常同时包含多个动作、对象和数值要求。查询粒度过大时，关键违规句容易被其他内容稀释；查询粒度过小时，又可能丢失章节语境和适用场景。因此，需要同时保留父级上下文并提高具体要求的证据召回能力。")
    doc.add_heading("2.3 复杂审查任务的自主决策、工具核验与异步并行协同难题（成员A主笔，成员B协作）", level=2)
    body(doc, "规程审查不仅包含法规合规判断，还涉及错别字、重复内容、适用性、数值阈值和争议复核等任务。不同任务的处理路径、工具和计算开销存在差异，串行执行效率较低。尤其对于“不得超过”“低于后方可恢复”等方向性阈值，纯语言模型容易混淆上下限或忽略单位差异；对于模型分歧、数值冲突和低置信结果，又不能强制输出结论。因此，需要由主智能体理解任务、调用确定性工具、并行调度专业智能体，并自主决定自动裁决或升级人工。")
    doc.add_heading("2.4 多来源审查经验的分层沉淀与持续复用难题（成员A/B共同撰写）", level=2)
    body(doc, "审查过程中产生的领域规则、稳定流程、主智能体裁决和人工修改意见具有不同可信度、稳定性和复用方式。若仅保存最终结果，系统无法复用有效经验，也难以避免重复错误。因此，需要建立覆盖提示词、Skill、智能体裁决和人工反馈数据飞轮的多层经验提炼机制。")
    table(doc, ["创新点", "主要方法", "预期作用"], [
        ["多结构规则知识库", "章节感知切分 + 父级上下文 + 适用性标签", "统一组织异构规则文档"],
        ["结构感知多查询检索", "父块语义保底 + 片段补充 + 混合检索重排", "兼顾上下文与精准召回"],
        ["主智能体、工具核验与异步并行审查", "任务理解 + 数值工具 + 并行调度 + 争议升级", "提高审查效率与可信度"],
        ["多层经验持续进化", "提示词 + Skill + 智能体裁决 + 人工反馈飞轮", "沉淀并复用审查经验"],
    ], [3.8, 7.0, 5.0])

    # Chapter 3
    doc.add_heading("第三章 方案论证与设计（成员A统稿，成员B/C协作）", level=1)
    doc.add_heading("3.1 多结构规则知识库构建（成员B主笔）", level=2)
    body(doc, "系统面向法规、规程、标准和实施细则等不同结构的规则文档进行章节感知切分，保留篇、章、节、条以及父级上下文等结构信息，并记录来源文档和适用场景。对于突出矿井专用条款，系统依据文档与章节结构添加适用性标签，使知识库不仅保存文本内容，也保存法规层级和应用边界。")
    placeholder(doc, "待补图：异构规则文档—章节结构识别—结构化切分—适用性标注—多结构规则知识库")
    caption(doc, "图3-1  多结构规则知识库构建流程")
    doc.add_heading("3.2 待审文档结构感知切分与多查询检索（成员B主笔）", level=2)
    body(doc, "待审文档按照标题层级和内容长度切分，保留父级章节上下文和原文位置映射。检索阶段以待审父块整体查询作为语义保底，同时将块内编号项和关键句切分为查询片段，用于补充召回容易被长文本稀释的具体动作、对象和数值要求。各查询结果经 BGE-M3 稠密与稀疏召回、RRF 融合、重排序和证据去重后进入审查链。本文将该方法称为“结构感知父块保底与片段补充检索”，不将其夸大为已实现的完整父子向量索引。")
    doc.add_heading("3.3 主智能体自主决策、工具核验与多智能体异步并行审查（成员A主笔，成员B协作）", level=2)
    body(doc, "主智能体负责理解任务、读取技能说明、选择工具和调度审查流程。审查引擎先批量完成检索预计算，再并行执行合规审查、错别字检查和文档级重复检查；块级任务也可并发执行。合规链内部根据初审结果按需调用数值核验、二次验证和立场分类。初审与核验分歧、数值冲突、低置信结论和调用异常将自动进入升级队列，由主智能体继续裁决或升级人工。")
    body(doc, "对于数值阈值问题，系统采用“语言模型抽取、确定性代码比较”的工具核验方式。语言模型从待审内容和法规依据中识别参数对象、数值、单位、比较方向与适用场景；数值工具完成单位归一化、上下限方向编码和严格程度比较，并将工具结论注入二次验证智能体。当初审判断合规但内容包含报警、断电、复电、停工等阈值语义时，系统执行反向核验，发现疑似超限或方向错误后自动进入升级队列。")
    placeholder(doc, "待补图：主智能体策略—专业智能体异步并行—数值等工具核验—冲突检测—升级队列与人工裁决")
    caption(doc, "图3-2  主智能体自主决策、工具核验与异步并行审查流程")
    doc.add_heading("3.4 多层经验保存、提炼与持续进化机制（成员A/B共同撰写）", level=2)
    body(doc, "系统采用多层方式保存和复用经验：第一层将稳定的领域约束和判断规则沉淀为提示词；第二层将可重复执行的切分、审查和升级流程封装为 Skill；第三层将主智能体的争议裁决保存为智能体标注；第四层将人工接受、驳回和改写意见作为最高优先级经验写入数据飞轮。不同层级经验分别服务于即时推理、流程复用、错误复盘和评测集扩充，形成从运行反馈到系统优化的持续进化路径。")

    # Chapter 4
    doc.add_heading("第四章 原理分析及硬件电路图（成员C主笔，成员A/B协作）", level=1)
    doc.add_heading("4.1 软件系统研发架构（成员A/C共同撰写）", level=2)
    body(doc, "软件系统采用分层、解耦和异步处理的研发架构。交互层包括 Web 审查客户端与微信小程序；业务服务层通过 FastAPI 提供统一接口和实时消息推送；任务数据层利用 SQLite 队列连接业务服务与常驻 worker；智能处理层由 LangChain 工具组件、LangGraph 智能体编排、本地混合检索、专业审查智能体和 Word 修订工具组成；经验层负责保存 Prompt、Skill、智能体裁决和人工反馈数据。")
    body(doc, "该架构将高频业务请求与耗时智能任务分离，使客户端、后端服务、检索模型和专业智能体可以独立开发与扩展。规则知识库构建、待审文档检索、数值核验、异步并行审查和经验沉淀等关键原理已在第三章展开，本节重点说明各模块之间的工程连接关系。")
    placeholder(doc, "待补图：Web/微信小程序—FastAPI业务服务—SQLite任务队列—LangChain/LangGraph智能处理—知识库与经验层")
    caption(doc, "图4-1  软件系统研发架构")
    doc.add_heading("4.2 硬件设计及移动终端原理图（成员C主笔）", level=2)
    body(doc, "本作品不额外设计外接传感器电路，而是复用智能手机内置的麦克风、摄像头、定位、触摸屏、本地存储和无线通信模块。微信小程序通过终端接口完成现场信息采集、任务交互和结果展示，并通过 Wi-Fi、4G 或 5G 与智能服务端通信。该方案能够在五天开发周期内完成稳定原型，同时避免将小程序本身错误表述为自主设计的嵌入式硬件。")
    placeholder(doc, "待补图：麦克风/摄像头/定位/存储/触摸屏—微信小程序—无线通信—智能服务端")
    caption(doc, "图4-2  智能手机移动终端硬件功能架构")

    # Chapter 5
    doc.add_heading("第五章 软件设计与流程（成员C统稿，成员A/B分项撰写）", level=1)
    doc.add_heading("5.1 平台技术架构与部署流程（成员C主笔，成员A协作）", level=2)
    body(doc, "平台采用前后端分离和任务异步处理架构。交互层由 React Web 审查客户端和计划开发的微信小程序组成；业务服务层采用 FastAPI 提供文档上传、状态查询、问题流、人工反馈和结果下载等接口；任务与数据层使用 SQLite 连接后端服务和常驻 worker；智能处理层由主智能体、v9 审查引擎、本地 BGE 检索模型、数值核验工具和 Word 文档适配器组成。")
    body(doc, "Web 客户端通过 HTTP 提交任务，并通过服务器发送事件接收实时进度与新增问题。FastAPI 后端不直接加载大型检索模型，而是将任务写入共享任务库；常驻 worker 轮询任务、复用已加载模型完成审查并持续回写结果，从而降低模型重复加载开销，并实现业务请求与耗时算法任务解耦。")
    table(doc, ["软件层次", "主要技术与模块", "核心职责"], [
        ["交互层", "React + Vite + Ant Design；微信小程序（预期）", "文档审查交互、现场问答与结果展示"],
        ["业务服务层", "FastAPI + HTTP/SSE 接口", "任务创建、状态流转、反馈接收与文件下载"],
        ["任务与数据层", "SQLite WAL 任务库", "跨进程通信、问题记录、升级项与标注存储"],
        ["智能处理层", "主智能体 + v9 worker + BGE + 审查工具", "任务决策、知识检索、并行审查与文档修订"],
    ], [3.1, 6.6, 6.1])
    placeholder(doc, "待补图：React Web/微信小程序—FastAPI—SQLite任务库—常驻worker—主智能体与审查工具")
    caption(doc, "图5-1  平台软件技术架构与部署流程")

    doc.add_heading("5.2 Web审查客户端设计与流程（成员C主笔）", level=2)
    body(doc, "Web 审查客户端基于 React、Vite 和 Ant Design 构建，采用三栏式审查界面展示任务状态、实时问题列表和 Word 段落内容。用户上传待审文档并选择矿井类型后，客户端创建审查任务，通过 SSE 长连接接收审查进度和新增问题；点击问题可定位并高亮对应原文位置。")
    body(doc, "对于每条问题，用户可以选择接受系统建议、驳回问题或输入自定义修改意见。接受或自定义修改将进入后台反馈队列，由主智能体生成最终替换文本并修改 Word 工作副本；客户端轮询反馈处理状态，刷新文档内容并支持下载审查后的文件。数据飞轮页面用于查看和管理人工反馈记录。")
    placeholder(doc, "待补图：Web审查客户端页面、问题定位高亮和人工反馈流程")
    caption(doc, "图5-2  Web审查客户端设计与交互流程")

    doc.add_heading("5.3 微信小程序客户端设计与流程（成员C主笔，预期实现）", level=2)
    body(doc, "微信小程序作为移动端现场交互入口，计划设置首页、风险问答、现场上报、处置确认和历史记录等功能模块。小程序调用智能手机已有的麦克风和摄像头能力采集语音、文字与图片，通过 HTTPS 接口向服务端创建现场研判任务，并以对话卡片展示智能体追问、法规依据、风险等级和处置建议。")
    body(doc, "小程序首先提交用户问题和现场信息；当主智能体判断信息不足时，服务端返回结构化追问，小程序引导用户继续补充；当风险结论形成后，用户确认或修正处置建议，最终结果写入历史记录和反馈数据。考虑比赛时间限制，小程序优先完成文字问答、拍照上传、结果展示和历史记录，语音输入作为可选增强功能。")
    placeholder(doc, "待补图：微信小程序页面结构、现场问答时序与风险结果卡片")
    caption(doc, "图5-3  微信小程序客户端设计与交互流程")

    doc.add_heading("5.4 业务后端设计与流程（成员C主笔）", level=2)
    body(doc, "业务后端采用 Python FastAPI 框架开发，负责连接客户端与智能处理任务。现有 v9 Web 接口包括文档上传、任务状态、段落模型、问题增量流、人工反馈、反馈状态、数据飞轮和修订文档下载。后端通过 SQLite 任务库与 worker 通信，不直接导入审查引擎，从而避免 Web 服务重载导致模型反复加载。")
    body(doc, "文档上传后，后端保存文件并创建 pending 任务；worker 领取任务后依次更新 parsing、reviewing 和 done 等状态。审查过程中，后端通过 SSE 将任务进度和新增问题增量推送给客户端。计划为微信小程序增加现场任务创建、多模态文件上传、追问回复、风险结果查询和处置确认接口，并复用现有任务状态与反馈机制。")
    placeholder(doc, "待补图：客户端请求—FastAPI接口—任务库—worker处理—SSE/结果返回时序")
    caption(doc, "图5-4  业务后端接口与异步任务流程")

    doc.add_heading("5.5 智能体与算法后端设计及流程（成员A/B共同撰写，部分预期实现）", level=2)
    body(doc, "智能体与算法后端采用 Python、LangChain 和 LangGraph 构建统一智能体框架。LangChain 用于封装大语言模型、提示词模板、法规检索器、数值核验、文档处理和结果存储等工具；LangGraph 用于表达主智能体与专业智能体之间的状态流转、条件分支、并行节点、异常重试和人工介入节点。常驻 worker 负责加载本地 BGE-M3 与重排序模型，并消费异步审查任务。")
    body(doc, "主智能体首先理解用户任务，生成包含目标、执行步骤、智能体分工、所需工具和升级条件的审查策略，再通过 LangGraph 调度知识检索智能体、合规审查智能体、错别字智能体、重复性检查智能体、数值核验智能体和结果修订智能体。各专业智能体在独立上下文中执行任务，并将结构化结果返回主智能体；主智能体根据置信度、工具结论和冲突情况决定合并结果、再次检索或升级人工。（完整 LangGraph 统一编排与策略生成节点为预期完善方案）")
    body(doc, "算法处理流程将检索预计算与语言模型审查分离，通过批量检索、专业智能体异步并行和块级并发降低总审查时间。人工接受或自定义修改后，结果修订智能体生成最终替换文本，并由 Word 适配器定位和修改对应段落。")
    placeholder(doc, "待补图：LangChain工具层—LangGraph主智能体策略生成—专业智能体异步并行—冲突裁决—人工节点—Word修订")
    caption(doc, "图5-5  基于LangChain与LangGraph的智能体算法后端流程（部分预期实现）")

    doc.add_heading("5.6 数据库与经验存储设计及流程（成员C主笔，成员A/B协作）", level=2)
    body(doc, "平台当前采用 SQLite WAL 模式保存跨进程任务数据。jobs 表记录文档路径、矿井类型、任务状态、审查进度和段落模型位置；issues 表记录问题类型、原文位置、法规依据、修改建议、人工动作和主智能体修改状态；escalations 表保存需要主智能体或人工处理的争议项；annotations 表保存智能体裁决和人工反馈形成的经验数据。")
    body(doc, "数据库同时承担运行状态保存和经验沉淀职责。运行数据支持任务断点恢复、实时问题流和反馈处理；经验数据可按来源区分智能体与人工意见，并导出为评测案例。提示词和 Skill 文件作为文件级经验保存，SQLite 标注作为实例级经验保存，二者共同支持系统持续优化。计划中的小程序任务和处置记录可复用现有任务、问题和标注模型，并按任务来源增加移动端标识。")
    table(doc, ["数据对象", "主要内容", "用途"], [
        ["jobs", "任务状态、进度、文件及段落模型路径", "异步任务调度与状态恢复"],
        ["issues", "问题位置、依据、建议及人工反馈", "实时展示、裁决与Word修改"],
        ["escalations", "分歧、冲突、低置信和异常项", "主智能体自主裁决与人工升级"],
        ["annotations", "智能体裁决、人工接受/驳回/改写", "数据飞轮、误差分析与评测集扩充"],
        ["Prompt/Skill文件", "稳定领域规则和可复用流程", "提示约束与主智能体流程调用"],
    ], [3.0, 7.1, 5.7])
    placeholder(doc, "待补图：jobs/issues/escalations/annotations关系及Prompt、Skill经验存储层次")
    caption(doc, "图5-6  数据库与多层经验存储结构")

    # Chapter 6
    doc.add_heading("第六章 系统测试与分析（成员B统稿，成员A/C分项测试）", level=1)
    body(doc, "本章以人工复核的统一黄金标注集和典型规程审查任务作为测试基础，从检索效果、专业智能体功能、主智能体策略生成时间、多智能体异步并行性能以及完整系统流程五个方面验证系统。对于尚未完成的测试，正文中标注为预期测试，并在项目完成后补充具体结果。")

    doc.add_heading("6.1 多结构知识库与待审文档检索效果测试（成员B主笔）", level=2)
    body(doc, "检索效果测试用于验证不同知识组织和查询方式对法规证据召回质量的影响。实验在统一黄金证据组上，对比传统稠密检索、稠密与稀疏混合检索、父块整体查询以及父块保底与片段补充检索。每种方法使用相同知识库、测试案例和候选数量，报告证据覆盖率、Hit@K、Recall@K、MRR、NDCG@K 和平均检索时间。")
    table(doc, ["对比方案", "主要设置", "核心测试指标"], [
        ["稠密检索基线", "待审父块直接进行Dense召回", "Hit@K、Recall@K、平均耗时"],
        ["混合检索", "Dense + Learned Sparse + RRF", "覆盖率、MRR、NDCG@K"],
        ["混合检索与重排序", "混合召回 + BGE Reranker", "前排命中率、NDCG@K"],
        ["父块保底与片段补充", "父块查询 + 编号项/关键句查询 + 去重", "覆盖率、Recall@K、漏召回案例"],
    ], [4.0, 6.4, 5.4])
    placeholder(doc, "待补图表：不同检索方案的Hit@K、Recall@K、NDCG和检索耗时对比")
    caption(doc, "图6-1  多种检索方案效果对比")

    doc.add_heading("6.2 各专业智能体功能测试（成员A/B共同撰写）", level=2)
    body(doc, "专业智能体功能测试分别验证各智能体能否完成其职责，并分析错误类型。测试集覆盖合规问题、错别字、重复内容、数值阈值冲突、法规适用性争议和文档修改任务。每个智能体独立运行，使用人工标注结果评价其输出，报告准确率、召回率、F1值、任务成功率和典型失败案例。")
    table(doc, ["测试对象", "主要功能", "核心指标", "当前状态"], [
        ["知识检索智能体", "生成查询并召回法规证据", "证据命中率、Recall@K", "已实现基础能力，待独立测试"],
        ["合规审查智能体", "依据法规判断不合规内容", "Precision、Recall、F1", "已实现，待统一集测试"],
        ["错别字检查智能体", "识别错别字和不规范用词", "准确率、召回率、误报率", "已实现，待测试"],
        ["重复性检查智能体", "发现文档内重复或高度相似内容", "准确率、召回率、平均耗时", "已实现，待测试"],
        ["数值核验智能体", "抽取数值并调用确定性比较工具", "数值判断准确率、漏报率", "已实现，待对照测试"],
        ["结果修订智能体", "根据建议和人工意见修改Word", "定位成功率、修改成功率", "已实现闭环，待批量测试"],
        ["现场问答智能体", "追问现场信息并给出处置建议", "任务完成率、依据正确率", "预期实现/待测试"],
    ], [3.3, 5.0, 4.1, 3.4])
    placeholder(doc, "待补图表：各专业智能体准确率、召回率、任务成功率及典型案例")
    caption(doc, "图6-2  各专业智能体功能测试结果")

    doc.add_heading("6.3 主智能体策略生成时间与决策质量测试（成员A主笔，成员B协作）", level=2)
    body(doc, "主智能体策略生成测试用于评价系统接收任务后生成执行方案的效率和合理性。测试任务分为规则文档入库、待审文档审查、争议项复核、Word修订和现场风险问答等类型。记录从接收任务到输出完整策略所需时间，并检查策略是否正确识别任务类型、选择合适智能体与工具、生成合理执行顺序和设置必要升级条件。（统一 LangGraph 策略生成节点为预期实现/待测试）")
    table(doc, ["测试指标", "计算方式", "测试目的"], [
        ["策略生成响应时间", "任务输入至完整策略输出的时间", "验证主智能体决策效率"],
        ["任务识别准确率", "正确识别任务类型的案例占比", "验证任务理解能力"],
        ["工具选择正确率", "正确选择必要工具的比例", "验证工具调用规划能力"],
        ["策略执行匹配度", "实际执行步骤与人工参考策略的匹配程度", "验证策略合理性"],
        ["升级决策准确率", "应升级案例与实际升级案例的一致程度", "验证风险控制能力"],
    ], [4.0, 7.0, 4.8])
    placeholder(doc, "待补图表：不同任务类型的策略生成时间、工具选择正确率和策略匹配度")
    caption(doc, "图6-3  主智能体策略生成时间与决策质量测试")

    doc.add_heading("6.4 多智能体异步并行审查性能测试（成员A/B共同撰写）", level=2)
    body(doc, "异步并行性能测试对比串行审查、专业智能体并行审查以及块级并发审查三种模式。在相同文档、模型和接口条件下，记录总审查时间、首条问题返回时间、平均块处理时间、吞吐量、接口异常率和最终审查结果一致性，用于验证并行架构是否在不明显降低质量的情况下提升处理效率。")
    table(doc, ["对比模式", "执行方式", "核心指标"], [
        ["串行基线", "检索、合规、错别字、重复检查依次执行", "总耗时、结果质量"],
        ["专业智能体并行", "合规链、错别字链、重复性链并行", "加速比、首条结果时间"],
        ["专业智能体并行+块级并发", "多链并行并同时处理多个待审块", "吞吐量、异常率、加速比"],
    ], [4.2, 6.6, 5.0])
    placeholder(doc, "待补图表：不同并行模式的总审查时间、首条结果时间、吞吐量和加速比")
    caption(doc, "图6-4  多智能体异步并行审查性能对比")

    doc.add_heading("6.5 Web与微信小程序端到端系统测试（成员C主笔）", level=2)
    body(doc, "端到端系统测试验证用户从提交任务到获得结果并反馈修订的完整流程。Web端重点测试文档上传、实时问题推送、原文定位、人工反馈、Word修改和结果下载；微信小程序重点测试文字提问、图片上传、追问交互、结果展示、历史记录和弱网重试。（微信小程序相关测试为预期实现/待测试）")
    table(doc, ["测试场景", "主要检查内容", "核心指标"], [
        ["Web文档审查闭环", "上传—审查—定位—反馈—修改—下载", "流程成功率、端到端耗时"],
        ["Word修改", "段落定位、连续修改、修改后下载", "定位成功率、修改成功率"],
        ["微信小程序问答", "提问、追问、法规依据与结果展示", "任务完成率、响应时间"],
        ["图片上传与现场上报", "拍照上传、内容接收与记录保存", "上传成功率、处理时间"],
        ["弱网与异常恢复", "超时、重试、任务恢复与历史记录", "恢复成功率、恢复时间"],
    ], [4.0, 7.1, 4.7])
    placeholder(doc, "待补图表：Web端与微信小程序完整流程测试结果及界面截图")
    caption(doc, "图6-5  Web与微信小程序端到端系统测试")

    # Chapter 7
    doc.add_heading("第七章 总结（成员A主笔，全员校核）", level=1)
    body(doc, "本作品围绕煤矿规程审查和现场风险研判，设计了具备任务理解、工具调用、自主升级和人工协同能力的移动端协同多智能体系统。当前项目已完成结构化切分、本地混合检索、多链并行审查、数值工具核验、升级队列、主智能体裁决、人工反馈、Word 修订和数据飞轮等核心链路。")
    body(doc, "下一阶段将重点完成微信小程序，补齐现场风险对话与记录保存工具，并在统一黄金集上开展算法、智能体和系统工程测试。通过将作业前规程审查、作业中现场研判和作业后反馈沉淀连接起来，系统有望为煤矿安全管理提供更高效、可信和可追溯的技术支撑。")

    doc.add_heading("研究成果及应用价值（成员A主笔，全员提供材料）", level=1)
    item(doc, "（1）形成一套面向煤矿规程的结构保持型切分与本地混合检索方法。")
    item(doc, "（2）形成“语言模型抽取、确定性代码比较”的数值核验工具，并支持反向漏报核验。")
    item(doc, "（3）形成争议升级、主智能体裁决、人工反馈、Word 修订和标注回流的闭环系统。")
    item(doc, "（4）拟形成复用智能手机硬件能力的微信小程序现场风险研判原型。")
    body(doc, "该系统可应用于煤矿作业规程审查、安全培训、现场隐患上报和风险处置，也可迁移至电力、化工和制造业等强规范、重安全场景。竞赛原型主要用于验证技术可行性，实际下井部署仍需满足矿用防爆、本安认证和生产通信网络等要求。")

    doc.add_heading("参考文献（成员B整理，成员A终审）", level=1)
    refs = [
        "[1] 国家矿山安全监察局. 煤矿安全规程（2025版）.",
        "[2] Lewis P, Perez E, Piktus A, et al. Retrieval-Augmented Generation for Knowledge-Intensive NLP Tasks. NeurIPS, 2020.",
        "[3] Chen J, et al. BGE M3-Embedding: Multi-Lingual, Multi-Functionality, Multi-Granularity Text Embeddings Through Self-Knowledge Distillation. 2024.",
        "[4] Yao S, Zhao J, Yu D, et al. ReAct: Synergizing Reasoning and Acting in Language Models. ICLR, 2023.",
        "[5] 其余煤矿法规、智能体、边缘计算和多模态相关文献由成员B统一补充并按竞赛格式整理。",
    ]
    for ref in refs:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.74)
        p.paragraph_format.first_line_indent = Cm(-0.74)
        p.paragraph_format.line_spacing = 1.25
        font(p.add_run(ref), "宋体", 10.5)

    # Footer
    for sec in doc.sections:
        p = sec.footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        font(p.add_run("研电赛技术论文简要初稿（含分工标注）"), "宋体", 9, color="777777")

    doc.core_properties.title = "研电赛技术论文简要初稿（含三人分工标注）"
    doc.core_properties.subject = "煤矿安全规程审查与现场风险研判多智能体系统"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
