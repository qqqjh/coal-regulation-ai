from __future__ import annotations

import io
import json
import re
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml.ns import qn


ROOT = Path(r"D:\work\project\coal-regulation-ai")
OUT = ROOT / "\u5206\u5de5_\u4e09\u4eba\u6574\u5408\u7248_xml.docx"

SRC_MINE = Path(r"D:\work\项目\研电赛\分工_第二三章和第六章1-4.docx")
SRC_LZY = Path(r"D:\work\项目\研电赛\分工_lzy.docx")
SRC_YY = Path(r"D:\work\项目\研电赛\分工_yy.docx")

R_ATTRS = [
    qn("r:embed"),
    qn("r:id"),
    qn("r:link"),
]


def block_text(el) -> str:
    if el.tag != qn("w:p"):
        return ""
    return "".join(t.text or "" for t in el.iter(qn("w:t"))).strip()


def body_blocks(doc: Document):
    return [el for el in doc.element.body if el.tag != qn("w:sectPr")]


def find_block_index(doc: Document, prefix: str) -> int:
    for idx, el in enumerate(body_blocks(doc)):
        text = block_text(el)
        if text.startswith(prefix):
            return idx
    raise RuntimeError(f"heading not found: {prefix}")


def remap_relationships(el, source: Document, target: Document) -> None:
    rid_map: dict[str, str] = {}
    for node in el.iter():
        for attr in R_ATTRS:
            rid = node.get(attr)
            if not rid or rid in rid_map:
                continue
            rel = source.part.rels.get(rid)
            if rel is None:
                continue
            if rel.is_external:
                new_rid = target.part.relate_to(rel.target_ref, rel.reltype, is_external=True)
            elif rel.reltype == RT.IMAGE:
                image_part = source.part.related_parts[rid]
                new_rid, _ = target.part.get_or_add_image(io.BytesIO(image_part.blob))
            else:
                target_part = source.part.related_parts.get(rid)
                if target_part is None:
                    continue
                new_rid = target.part.relate_to(target_part, rel.reltype)
            rid_map[rid] = new_rid
    for node in el.iter():
        for attr in R_ATTRS:
            rid = node.get(attr)
            if rid in rid_map:
                node.set(attr, rid_map[rid])


def clone_section(source: Document, target: Document, start_prefix: str, end_prefix: str):
    blocks = body_blocks(source)
    start = find_block_index(source, start_prefix)
    end = find_block_index(source, end_prefix)
    clones = []
    for el in blocks[start:end]:
        copied = deepcopy(el)
        remap_relationships(copied, source, target)
        clones.append(copied)
    return clones


def replace_section(target: Document, source: Document, start_prefix: str, end_prefix: str) -> int:
    body = target.element.body
    blocks = body_blocks(target)
    start = find_block_index(target, start_prefix)
    end = find_block_index(target, end_prefix)
    old_blocks = blocks[start:end]
    insert_at = body.index(old_blocks[0])
    new_blocks = clone_section(source, target, start_prefix, end_prefix)
    for el in old_blocks:
        body.remove(el)
    for offset, el in enumerate(new_blocks):
        body.insert(insert_at + offset, el)
    return len(new_blocks)


def iter_text_nodes(doc: Document):
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            yield run
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        yield run


def replace_inline_text(doc: Document, replacements: list[tuple[str, str]]) -> dict[str, int]:
    counts: dict[str, int] = {}
    for run in iter_text_nodes(doc):
        if not run.text:
            continue
        text = run.text
        for old, new in replacements:
            if old in text:
                c = text.count(old)
                text = text.replace(old, new)
                counts[old] = counts.get(old, 0) + c
        run.text = text
    return counts


def main() -> None:
    target = Document(str(SRC_YY))
    mine = Document(str(SRC_MINE))
    lzy = Document(str(SRC_LZY))

    steps = [
        ("mine", mine, "\u7b2c\u4e8c\u7ae0", "\u7b2c\u4e09\u7ae0"),
        ("mine", mine, "\u7b2c\u4e09\u7ae0", "\u7b2c\u56db\u7ae0"),
        ("lzy", lzy, "4.1", "4.2"),
        ("lzy", lzy, "5.2", "5.3"),
        ("lzy", lzy, "5.4", "5.5"),
        ("lzy", lzy, "5.5", "5.6"),
        ("mine", mine, "6.1", "6.5"),
    ]
    step_counts = []
    for source_name, source, start, end in steps:
        count = replace_section(target, source, start, end)
        step_counts.append({"source": source_name, "start": start, "end": end, "blocks": count})

    replacements = [
        ("\u672c\u5730SQLite\u6570\u636e\u5e93", "MongoDB\u6570\u636e\u5e93"),
        ("SQLite \u4efb\u52a1\u5e93", "MongoDB \u4efb\u52a1\u96c6\u5408"),
        ("SQLite\u4efb\u52a1\u5e93", "MongoDB\u4efb\u52a1\u96c6\u5408"),
        ("SQLite \u4efb\u52a1\u6570\u636e\u5e93", "MongoDB \u4efb\u52a1\u6570\u636e\u5e93"),
        ("SQLite\u4efb\u52a1\u6570\u636e\u5e93", "MongoDB\u4efb\u52a1\u6570\u636e\u5e93"),
        ("\u6570\u636e\u5e93\u91c7\u7528SQLite", "\u6570\u636e\u5e93\u91c7\u7528MongoDB"),
        ("\u91c7\u7528SQLite", "\u91c7\u7528MongoDB"),
        ("Chroma\u5411\u91cf\u5e93", "Milvus\u5411\u91cf\u5e93"),
        ("Chroma \u5411\u91cf\u5e93", "Milvus \u5411\u91cf\u5e93"),
    ]
    replacement_counts = replace_inline_text(target, replacements)

    target.save(str(OUT))
    print(json.dumps({"output": str(OUT), "steps": step_counts, "replacements": replacement_counts}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
