from __future__ import annotations

import hashlib
import json
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


SRC = Path("D:/work") / "\u9879\u76ee" / "\u7814\u7535\u8d5b" / "\u7814\u7535\u8d5b\u6280\u672f\u8bba\u6587\u6700\u7ec8\u7248v3.docx"
OUT_WORKSPACE = Path("D:/work/project/coal-regulation-ai") / "\u7814\u7535\u8d5b\u6280\u672f\u8bba\u6587\u6700\u7ec8\u7248v4.docx"


P56 = [
    "\u6570\u636e\u5c42\u91c7\u7528 MongoDB + Milvus \u534f\u540c\u65b9\u6848\uff0cMongoDB \u4fdd\u5b58\u77e5\u8bc6\u5e93\u3001\u6587\u6863\u3001\u5207\u5206\u7247\u6bb5\u3001\u5ba1\u67e5\u4efb\u52a1\u3001\u95ee\u9898\u6e05\u5355\u3001\u4eba\u5de5\u53cd\u9988\u3001\u73b0\u573a\u4e0a\u62a5\u8bb0\u5f55\u548c\u6a21\u578b\u8c03\u7528\u65e5\u5fd7\uff1bMilvus \u4fdd\u5b58\u6cd5\u89c4\u6761\u6b3e\u3001\u89c4\u7a0b\u7247\u6bb5\u548c\u7ecf\u9a8c\u6837\u672c\u7684\u5411\u91cf\u7d22\u5f15\u3002\u4e1a\u52a1\u6570\u636e\u4fdd\u8bc1\u4efb\u52a1\u8fc7\u7a0b\u53ef\u8ffd\u6eaf\uff0c\u7ecf\u9a8c\u6570\u636e\u5219\u6309\u6765\u6e90\u548c\u4f5c\u7528\u5206\u7ea7\u4fdd\u5b58\u3002",
    "\u7ecf\u9a8c\u673a\u5236\u5206\u4e3a\u9884\u7f6e\u7ecf\u9a8c\u3001\u8fd0\u884c\u7ecf\u9a8c\u548c\u53cd\u9988\u7ecf\u9a8c\u4e09\u7ea7\u3002\u9884\u7f6e\u7ecf\u9a8c\u5305\u62ec\u8d5b\u524d\u4eba\u5de5\u7f16\u5199\u7684 Prompt \u6a21\u677f\u3001Skill \u6d41\u7a0b\u3001\u5ba1\u67e5\u89c4\u5219\u548c\u5de5\u5177\u8c03\u7528\u89c4\u8303\uff0c\u5b83\u4eec\u662f\u7cfb\u7edf\u521d\u59cb\u80fd\u529b\uff0c\u4e0d\u7531\u6570\u636e\u98de\u8f6e\u81ea\u52a8\u751f\u6210\uff1b\u8fd0\u884c\u7ecf\u9a8c\u4fdd\u5b58\u4e3b\u667a\u80fd\u4f53\u7684\u4efb\u52a1\u5224\u65ad\u3001\u8bc1\u636e\u9009\u62e9\u3001\u5de5\u5177\u8c03\u7528\u548c\u51b2\u7a81\u88c1\u51b3\u8fc7\u7a0b\uff0c\u7528\u4e8e\u8ffd\u6eaf\u3001\u8bc4\u6d4b\u548c\u6392\u9519\uff1b\u53cd\u9988\u7ecf\u9a8c\u6765\u81ea\u5ba1\u67e5\u4eba\u5458\u7684\u63a5\u53d7\u3001\u9a73\u56de\u3001\u4fee\u6539\u3001\u8865\u5145\u8bc1\u636e\u548c\u73b0\u573a\u5904\u7f6e\u786e\u8ba4\u3002\u6570\u636e\u98de\u8f6e\u53ea\u56f4\u7ed5\u4eba\u5de5\u53cd\u9988\u95ed\u73af\u5f00\u5c55\uff1a\u7cfb\u7edf\u5c06\u53cd\u9988\u8bb0\u5f55\u6c89\u6dc0\u4e3a\u8bc4\u6d4b\u6837\u672c\u3001\u95ee\u9898\u5206\u5e03\u7edf\u8ba1\u548c\u4eba\u5de5\u7ef4\u62a4\u9884\u7f6e Prompt/Skill \u7684\u4f9d\u636e\uff0c\u4f46\u4e0d\u81ea\u52a8\u751f\u6210\u6216\u66ff\u6362 Prompt/Skill\u3002",
]

P62_AFTER = "\u88686-2\u4ec5\u4fdd\u7559\u4e09\u4e2a\u5df2\u8bc4\u4ef7\u667a\u80fd\u4f53\u3002\u5408\u89c4\u5ba1\u67e5\u667a\u80fd\u4f53\u4f7f\u7528\u4eba\u5de5\u9ec4\u91d1\u7ed3\u8bba\u8ba1\u7b97\u7ed3\u8bba\u5339\u914d\u51c6\u786e\u7387\u548c\u975e\u5408\u89c4 F1\uff0c\u5e76\u540c\u65f6\u62a5\u544a RAGAS \u6307\u6807\uff1b\u9519\u522b\u5b57\u68c0\u67e5\u548c\u91cd\u590d\u6027\u68c0\u67e5\u6682\u65e0\u4eba\u5de5\u91d1\u6807\uff0c\u56e0\u6b64\u4ec5\u62a5\u544a Faithfulness \u548c Relevancy\uff0c\u5f85\u8865\u5145\u6807\u6ce8\u540e\u518d\u8ba1\u7b97\u7ed3\u8bba\u7c7b\u6307\u6807\u3002"

TABLE62 = [
    ["\u6d4b\u8bd5\u5bf9\u8c61", "\u7ed3\u8bba\u5339\u914d\u51c6\u786e\u7387", "\u975e\u5408\u89c4 F1", "Faithfulness", "Relevancy"],
    ["\u5408\u89c4\u5ba1\u67e5\u667a\u80fd\u4f53", "56.82%", "26.67%", "0.1671", "0.4415"],
    ["\u9519\u522b\u5b57\u68c0\u67e5\u667a\u80fd\u4f53", "\u2014", "\u2014", "0.2581", "0.5083"],
    ["\u91cd\u590d\u6027\u68c0\u67e5\u667a\u80fd\u4f53", "\u2014", "\u2014", "0.3872", "0.5518"],
]


def set_para_text(paragraph, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def section_bounds(doc: Document) -> dict[str, tuple[int, int]]:
    heads: list[tuple[int, str]] = []
    for i, p in enumerate(doc.paragraphs):
        text = p.text.strip()
        if p.style.name.startswith("Heading") or re.match(r"^[1-7]\.\d+", text):
            heads.append((i, text))
    bounds = {}
    for idx, (start, text) in enumerate(heads):
        end = heads[idx + 1][0] if idx + 1 < len(heads) else len(doc.paragraphs)
        if text.startswith("5.6"):
            bounds["5.6"] = (start, end)
        if text.startswith("6.2"):
            bounds["6.2"] = (start, end)
    return bounds


def replace_body_paragraphs(doc: Document, key: str, paragraphs: list[str]) -> None:
    start, end = section_bounds(doc)[key]
    targets: list[int] = []
    for i in range(start + 1, end):
        text = doc.paragraphs[i].text.strip()
        if not text:
            continue
        if text.startswith(("\u56fe", "\u8868")):
            continue
        targets.append(i)
    if len(targets) < len(paragraphs):
        raise RuntimeError(f"not enough body paragraphs in {key}: {targets}")
    for para_idx, value in zip(targets, paragraphs):
        set_para_text(doc.paragraphs[para_idx], value)
    for para_idx in targets[len(paragraphs):]:
        set_para_text(doc.paragraphs[para_idx], "")


def set_cell(cell, text: str, header: bool = False) -> None:
    cell.text = text
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for p in cell.paragraphs:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = Pt(18)
        for r in p.runs:
            r.font.size = Pt(9)
            r.font.bold = header
            r.font.name = "Times New Roman"
            r._element.rPr.rFonts.set(qn("w:eastAsia"), "\u5b8b\u4f53")


def set_table_62(table) -> None:
    while len(table.columns) < 5:
        table.add_column(Cm(2.6))
    while len(table.rows) < 4:
        table.add_row()
    while len(table.rows) > 4:
        tr = table.rows[-1]._tr
        tr.getparent().remove(tr)
    for r_idx, row in enumerate(TABLE62):
        for c_idx, value in enumerate(row):
            set_cell(table.cell(r_idx, c_idx), value, header=(r_idx == 0))
    # If an old extra cell somehow remains, blank it deterministically.
    for r_idx in range(len(table.rows)):
        for c_idx in range(len(TABLE62[0]), len(table.rows[r_idx].cells)):
            set_cell(table.cell(r_idx, c_idx), "")


def count_chars(doc: Document):
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    text = "\n".join(parts)
    return len(re.findall(r"[\u4e00-\u9fff]", text)), len(re.sub(r"\s+", "", text)), text


def hash_table(table) -> str:
    return hashlib.sha256(table._tbl.xml.encode("utf-8")).hexdigest()


def main() -> None:
    doc = Document(str(SRC))
    table_hashes_before = [hash_table(t) for t in doc.tables]
    replace_body_paragraphs(doc, "5.6", P56)
    replace_body_paragraphs(doc, "6.2", [
        "\u4e13\u4e1a\u667a\u80fd\u4f53\u529f\u80fd\u6d4b\u8bd5\u7528\u4e8e\u9a8c\u8bc1\u5404\u667a\u80fd\u4f53\u662f\u5426\u5b8c\u6210\u804c\u8d23\u3002\u88686-2\u6309\u7ed3\u8bba\u5339\u914d\u51c6\u786e\u7387\u3001\u975e\u5408\u89c4 F1\u3001Faithfulness \u548c Relevancy \u7edf\u4e00\u5c55\u793a\u4e09\u4e2a\u5df2\u8bc4\u4ef7\u667a\u80fd\u4f53\u7684\u7ed3\u679c\u3002",
        P62_AFTER,
    ])
    set_table_62(doc.tables[2])
    doc.save(str(OUT_WORKSPACE))

    reopened = Document(str(OUT_WORKSPACE))
    cn, nonspace, text = count_chars(reopened)
    print(json.dumps({
        "output": str(OUT_WORKSPACE),
        "paragraphs": len(reopened.paragraphs),
        "tables": len(reopened.tables),
        "table_6_2_rows": len(reopened.tables[2].rows),
        "table_6_2_cols": len(reopened.tables[2].columns),
        "unchanged_table_hash_indexes": [
            i for i, (before, after) in enumerate(zip(table_hashes_before, [hash_table(t) for t in reopened.tables]))
            if before == after
        ],
        "cn_chars": cn,
        "nonspace_chars": nonspace,
        "data_flywheel_mentions": text.count("\u6570\u636e\u98de\u8f6e"),
        "auto_prompt_skill": "\u81ea\u52a8\u751f\u6210\u6216\u66ff\u6362 Prompt/Skill" in text,
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
