"""v9 Web 数据路径集成测试（不加载 GPU 引擎）

覆盖：上传建任务 → 段落模型落盘 → 写实时问题(带段落索引) → 人工反馈(custom)
     → 主智能体生成替换文本(真实LLM) → 改写工作副本 docx → 刷新段落模型。
验证"主智能体直接改 Word 对应位置"这条核心链路真正闭环。
"""
import os
import tempfile
from pathlib import Path

from docx import Document

import docx_adapter
import v9_worker
from review_queue import ReviewQueue


def main():
    work = Path(tempfile.mkdtemp())
    # 用独立的临时 DB / 工作目录，避免污染真实数据
    db = work / "test_queue.db"
    v9_worker.QUEUE_DB = db
    v9_worker.WORK_DIR = work / "webwork"
    v9_worker.WORK_DIR.mkdir(parents=True, exist_ok=True)

    # 1) 造一份带真实违规句的待审 docx
    doc = Document()
    doc.add_heading("第八章 安全技术措施", level=1)
    doc.add_paragraph("一、起吊作业")
    doc.add_paragraph("3、起吊设备前，选择支护锚索作为起吊点。")
    doc.add_paragraph("4、起吊时严禁大幅度斜拉或摆动。")
    docx_path = work / "pending.docx"
    doc.save(docx_path)

    q = ReviewQueue(db)

    # 2) 模拟后端 upload：建任务
    job_id = "testjob"
    q.create_job(job_id, "pending.docx", str(docx_path), "non_outburst")
    job = q.fetch_pending_job()
    assert job["status"] == "parsing"

    # 3) 模拟 worker 解析 + 落盘段落模型 + 定位问题段落
    parsed, chunks = docx_adapter.adapt(docx_path)
    para_path = work / "p.json"
    import json
    para_path.write_text(json.dumps(parsed, ensure_ascii=False), encoding="utf-8")
    q.update_job(job_id, status="reviewing", n_chunks=len(chunks),
                 paragraphs_path=str(para_path))

    original = "选择支护锚索作为起吊点"
    blocks = docx_adapter.locate_text_blocks(parsed, original)
    assert blocks, "应能定位违规原文所在段落"
    issue_id = q.add_issue(
        job_id, chunk_index=1, issue_type="compliance", status="不合规",
        block_indices=blocks, title="规则冲突", original_text=original,
        suggestion="改用专用吊装锚索作为起吊点，严禁使用支护锚索",
        regulation="GB/T 35056-2018 4.4.4.5",
        reason="将永久支护锚索用作起吊点违反禁止性规定",
    )
    print(f"问题已写入，定位段落 block={blocks}")

    # 4) 模拟后端 /feedback：人工 custom 反馈 + 写飞轮
    q.add_annotation("pending.docx", 1, original, "不合规", "human",
                     reason="人工确认违规并给出改法", extra={"action": "custom"})
    q.set_issue_feedback(issue_id, "custom",
                         "必须使用经计算校核的专用起吊锚杆，严禁使用支护锚索作为起吊点")

    # 5) worker 取反馈 → 主智能体生成替换文本 → 改写 docx（真实LLM）
    fb = q.fetch_pending_feedback()
    assert fb and fb["id"] == issue_id

    if not os.getenv("DASHSCOPE_API_KEY"):
        print("[跳过] 无 DASHSCOPE_API_KEY，跳过真实改写")
        return

    worker = v9_worker.V9Worker(load_engine=False)  # 反馈处理不需要引擎
    worker.queue = q  # 用同一临时库
    worker.process_feedback(fb)

    result = q.get_issue(issue_id)
    print(f"反馈处理结果 agent_applied={result['agent_applied']}")
    print(f"主智能体说明: {result['agent_note']}")
    assert result["agent_applied"] == 1, "应成功改写"

    # 6) 验证工作副本 docx 真的被改了，且原违规句已不在
    work_docx = v9_worker.WORK_DIR / f"{job_id}_working.docx"
    reparsed = docx_adapter.parse_docx(work_docx)
    full = "\n".join(b["text"] for b in reparsed["blocks"])
    print("\n改写后相关段落：")
    for b in reparsed["blocks"]:
        if "起吊点" in b["text"]:
            print("  ", b["text"])
    assert original not in full, "原违规句应已被替换"

    # 7) 验证飞轮记录了人工裁决（最高层）
    n = q.export_gold_cases(work / "gold.json", sources=("human",))
    assert n == 1, "人工反馈应进入飞轮"
    print("\n[OK] Web 数据路径闭环：建任务→段落定位→问题→人工反馈→主智能体改Word→飞轮")


if __name__ == "__main__":
    main()
