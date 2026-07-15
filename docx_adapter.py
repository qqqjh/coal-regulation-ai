"""docx 适配器 v1：Word 文档 ⇄ 前端段落模型 ⇄ v9 审查块（带段落索引）

为前后端打通而生。解决两个问题：
  1. 前端中间面板要按"段落"展示 Word，并支持点击问题 → 定位高亮某段。
  2. v9 引擎审查的"块"必须能映射回原 docx 的段落位置，主智能体才能精确改字。

核心做法：按 docx body 的真实顺序枚举 block（段落/表格），给每个 block 一个
稳定的全局 block_index；再把连续段落聚成 v9 审查块，每块记录 source_blocks
（来自哪些 block_index）。问题 → 块 → block_index → 段落高亮/改字，闭环成立。
"""
from __future__ import annotations

import re
import difflib
import unicodedata
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from docx import Document
from docx.document import Document as _DocxDocument
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from pending_chunk_packing_v9 import pack_pending_chunks_v9

# 章/节/条/编号标题检测（用于切块边界）
_HEADING_PATTERNS = [
    re.compile(r"^\s*第[一二三四五六七八九十百零\d]{1,4}章"),
    re.compile(r"^\s*第[一二三四五六七八九十百零\d]{1,4}节"),
    re.compile(r"^\s*第[一二三四五六七八九十百零\d]{1,5}条"),
    re.compile(r"^\s*[一二三四五六七八九十]{1,3}、"),
]
_CHAPTER_RE = re.compile(r"^\s*第[一二三四五六七八九十百零\d]{1,4}章")
_SECTION_RE = re.compile(r"^\s*(第[一二三四五六七八九十百零\d]{1,4}节|[一二三四五六七八九十]{1,3}、)")
_FORMAL_SECTION_RE = re.compile(r"^\s*第[一二三四五六七八九十百零\d]{1,4}节")
_ARTICLE_RE = re.compile(r"^\s*第[一二三四五六七八九十百零\d]{1,5}条")


def iter_block_items(document: _DocxDocument):
    """按文档真实顺序产出 (kind, obj)：kind ∈ {'paragraph','table'}。
    python-docx 的 doc.paragraphs/doc.tables 不保证相对顺序，必须走 body。"""
    body = document.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            yield "paragraph", Paragraph(child, document)
        elif child.tag == qn("w:tbl"):
            yield "table", Table(child, document)


def _table_to_rows(table: Table) -> List[List[str]]:
    rows = []
    for row in table.rows:
        rows.append([cell.text.strip() for cell in row.cells])
    return rows


def _is_heading(text: str) -> bool:
    return any(p.match(text) for p in _HEADING_PATTERNS)


def _is_structural_heading(text: str) -> bool:
    return bool(_CHAPTER_RE.match(text) or _FORMAL_SECTION_RE.match(text) or _ARTICLE_RE.match(text))


def parse_docx(path: str | Path) -> Dict[str, Any]:
    """解析 docx → 前端段落模型。

    返回：
      {
        "blocks": [
          {"block_index": 0, "kind": "paragraph"|"table",
           "text": str, "is_heading": bool, "style": str,
           "rows": [[...]]  # 仅 table},
          ...
        ],
        "n_blocks": int
      }
    """
    document = Document(str(path))
    blocks: List[Dict[str, Any]] = []
    for kind, obj in iter_block_items(document):
        if kind == "paragraph":
            text = obj.text or ""
            style = (obj.style.name if obj.style is not None else "") or ""
            blocks.append({
                "block_index": len(blocks),
                "kind": "paragraph",
                "text": text,
                "is_heading": _is_heading(text) or style.lower().startswith("heading"),
                "style": style,
            })
        else:  # table
            rows = _table_to_rows(obj)
            flat = "\n".join(" | ".join(r) for r in rows)
            blocks.append({
                "block_index": len(blocks),
                "kind": "table",
                "text": flat,
                "is_heading": False,
                "style": "Table",
                "rows": rows,
            })
    return {"blocks": blocks, "n_blocks": len(blocks)}


def build_chunks(parsed: Dict[str, Any]) -> List[Dict[str, Any]]:
    """段落模型 → v9 审查块。每块带 source_blocks（段落索引列表）。

    当前只在 Word block 层生成初始结构块，然后复用 pending_doc_chunking_v9
    抽出的通用 packing 规则：
      - 遇章/节/条标题 → 切出初始块
      - 表格单独成块并保留完整
      - 每个初始块携带 source_units，二次切分后仍能保留 source_blocks
      - 公共 packing 负责小标题合并、多级编号切分、相邻单元打包和超长兜底
    """
    blocks = parsed["blocks"]
    initial_chunks: List[Dict[str, Any]] = []
    cur_chapter = ""
    cur_section = ""
    cur_article = ""

    buf_texts: List[str] = []
    buf_indices: List[int] = []
    buf_units: List[Dict[str, Any]] = []

    def flush():
        nonlocal buf_texts, buf_indices, buf_units
        content = "\n".join(t for t in buf_texts if t.strip()).strip()
        if content:
            initial_chunks.append({
                "content": content,
                "chapter": cur_chapter,
                "section": cur_section,
                "article": cur_article,
                "page_range": "",
                "chunk_level": "paragraph",
                "char_count": len(content),
                "source_blocks": list(buf_indices),
                "source_units": list(buf_units),
            })
        buf_texts = []
        buf_indices = []
        buf_units = []

    for block in blocks:
        idx = block["block_index"]
        text = block["text"]

        if block["kind"] == "table":
            flush()
            content = f"【表格】\n{text}".strip()
            initial_chunks.append({
                "content": content,
                "chapter": cur_chapter,
                "section": cur_section,
                "article": cur_article,
                "page_range": "",
                "chunk_level": "table",
                "char_count": len(content),
                "oversized_table_chunk": True,
                "source_blocks": [idx],
                "source_units": [{"text": content, "source_blocks": [idx], "kind": "table"}],
            })
            continue

        if not text.strip():
            buf_indices.append(idx)  # 空段也归属当前块，保留索引连续性
            continue

        if _is_structural_heading(text) and buf_texts:
            flush()

        if _CHAPTER_RE.match(text):
            cur_chapter = text.strip()
            cur_section = ""
            cur_article = ""
        elif _SECTION_RE.match(text):
            cur_section = text.strip()
            cur_article = ""
        elif _ARTICLE_RE.match(text):
            cur_article = text.strip()

        buf_texts.append(text)
        buf_indices.append(idx)
        buf_units.append({"text": text, "source_blocks": [idx], "kind": block["kind"]})

    flush()
    return pack_pending_chunks_v9(initial_chunks)


def adapt(path: str | Path) -> Tuple[Dict[str, Any], List[Dict[str, Any]]]:
    """一步到位：返回 (前端段落模型, v9审查块列表)。"""
    parsed = parse_docx(path)
    chunks = build_chunks(parsed)
    return parsed, chunks


# ============ 改字：按 block_index 精确定位 docx 段落 ============

def _ordered_paragraph_map(document: _DocxDocument) -> Dict[int, Paragraph]:
    """block_index → Paragraph 对象（仅段落类型；表格 block 不在此 map）。"""
    mapping: Dict[int, Paragraph] = {}
    for block_index, (kind, obj) in enumerate(iter_block_items(document)):
        if kind == "paragraph":
            mapping[block_index] = obj
    return mapping


def _norm_text(text: str) -> str:
    return re.sub(r"\s+", "", unicodedata.normalize("NFKC", str(text or "")))


def _norm_span(full_text: str, target: str) -> Optional[Tuple[int, int]]:
    """返回 target 规范化后在 full_text 中对应的原始字符区间。"""
    norm_target = _norm_text(target)
    if not norm_target:
        return None
    norm_chars: List[str] = []
    index_map: List[int] = []
    for idx, ch in enumerate(str(full_text or "")):
        normalized = _norm_text(ch)
        if not normalized:
            continue
        for out_ch in normalized:
            norm_chars.append(out_ch)
            index_map.append(idx)
    norm_full = "".join(norm_chars)
    pos = norm_full.find(norm_target)
    if pos < 0:
        return None
    start = index_map[pos]
    end = index_map[pos + len(norm_target) - 1] + 1
    return start, end


def _replace_span_in_para(
    para: Paragraph,
    start: int,
    end: int,
    new_text: str,
    mark_color: Optional[Tuple[int, int, int]] = (0, 128, 0),
) -> bool:
    from docx.shared import RGBColor

    pos = 0
    changed = False
    for run in para.runs:
        run_start = pos
        run_end = pos + len(run.text)
        pos = run_end
        if run_end <= start or run_start >= end:
            continue
        prefix = run.text[: max(0, start - run_start)] if run_start <= start else ""
        suffix = run.text[max(0, end - run_start):] if run_end >= end else ""
        insert = new_text if not changed else ""
        run.text = prefix + insert + suffix
        if mark_color and insert:
            run.font.color.rgb = RGBColor(*mark_color)
        changed = True
    return changed


def _set_para_text(
    para: Paragraph,
    text: str,
    mark_color: Optional[Tuple[int, int, int]] = (0, 128, 0),
) -> bool:
    from docx.shared import RGBColor

    if para.runs:
        para.runs[0].text = text
        if mark_color:
            para.runs[0].font.color.rgb = RGBColor(*mark_color)
        for run in para.runs[1:]:
            run.text = ""
    else:
        run = para.add_run(text)
        if mark_color:
            run.font.color.rgb = RGBColor(*mark_color)
    return True


def _split_rewrite_lines(text: str, n: int) -> List[str]:
    lines = [line.strip() for line in re.split(r"\n+", str(text or "")) if line.strip()]
    if not lines:
        return ["" for _ in range(n)]
    if len(lines) >= n:
        return lines[: n - 1] + ["\n".join(lines[n - 1:])]
    return lines + ["" for _ in range(n - len(lines))]


def _extract_suggestion_pairs(suggestion: str) -> List[Tuple[str, str]]:
    """抽取“将 X 改/修改/调整为 Y”类局部替换建议。"""
    text = str(suggestion or "")
    if not text:
        return []
    pattern = re.compile(
        r"将\s*[\"'「『【]?\s*(.+?)\s*[\"'」』】]?\s*"
        r"(?:修改为|改为|更改为|调整为)\s*"
        r"[\"'「『【]?\s*(.+?)\s*[\"'」』】]?"
        r"(?=以|，|。|；|;|$)"
    )
    pairs = []
    for old, new in pattern.findall(text):
        old = old.strip(" \t\r\n\"'「」『』【】")
        new = new.strip(" \t\r\n\"'「」『』【】")
        if old and new and old != new:
            pairs.append((old, new))
            old_num = re.sub(r"^[<>≤≥=]+", "", old)
            new_num = re.sub(r"^[<>≤≥=]+", "", new)
            if old_num and new_num and old_num != old:
                pairs.append((old_num, new_num))
    return pairs


def apply_edit_at_block(document: _DocxDocument, block_index: int,
                        old_text: str, new_text: str,
                        mark_color: Optional[Tuple[int, int, int]] = (0, 128, 0)) -> bool:
    """在指定 block_index 的段落内做替换（主智能体改字用）。

    old_text 为空时整段替换为 new_text。返回是否成功。"""
    from docx.shared import RGBColor
    para_map = _ordered_paragraph_map(document)
    para = para_map.get(block_index)
    if para is None:
        return False

    if not old_text:
        # 整段替换：清空 runs，写入新文本到第一个 run（或新建）
        if para.runs:
            para.runs[0].text = new_text
            if mark_color:
                para.runs[0].font.color.rgb = RGBColor(*mark_color)
            for run in para.runs[1:]:
                run.text = ""
        else:
            run = para.add_run(new_text)
            if mark_color:
                run.font.color.rgb = RGBColor(*mark_color)
        return True

    if old_text not in para.text:
        span = _norm_span(para.text, old_text)
        if span is None:
            return False
        return _replace_span_in_para(para, span[0], span[1], new_text, mark_color)

    # 优先在单个 run 内替换（保留样式）
    for run in para.runs:
        if old_text in run.text:
            run.text = run.text.replace(old_text, new_text, 1)
            if mark_color:
                run.font.color.rgb = RGBColor(*mark_color)
            return True
    # 跨 run：把整段重写到第一个 run（会丢段内混合样式，但保证替换成功）
    replaced = para.text.replace(old_text, new_text, 1)
    if para.runs:
        para.runs[0].text = replaced
        if mark_color:
            para.runs[0].font.color.rgb = RGBColor(*mark_color)
        for run in para.runs[1:]:
            run.text = ""
        return True
    return False


def apply_suggestion_at_blocks(
    document: _DocxDocument,
    block_indices: List[int],
    suggestion: str,
    mark_color: Optional[Tuple[int, int, int]] = (0, 128, 0),
) -> bool:
    """按“将 X 改为 Y”建议在候选段落内做最小替换。"""
    pairs = _extract_suggestion_pairs(suggestion)
    if not pairs:
        return False
    para_map = _ordered_paragraph_map(document)
    indices = [i for i in block_indices if i in para_map]
    if not indices:
        indices = list(para_map.keys())
    for old, new in pairs:
        for bi in indices:
            if apply_edit_at_block(document, bi, old, new, mark_color):
                return True
    return False


def revert_suggestion_at_blocks(
    document: _DocxDocument,
    block_indices: List[int],
    suggestion: str,
    mark_color: Optional[Tuple[int, int, int]] = None,
) -> bool:
    """按“将 X 改为 Y”建议反向替换，用于撤回已改写裁决。"""
    pairs = _extract_suggestion_pairs(suggestion)
    if not pairs:
        return False
    para_map = _ordered_paragraph_map(document)
    indices = [i for i in block_indices if i in para_map]
    if not indices:
        indices = list(para_map.keys())
    for old, new in pairs:
        for bi in indices:
            if apply_edit_at_block(document, bi, new, old, mark_color):
                return True
    return False


def apply_rewrite_at_blocks(
    document: _DocxDocument,
    block_indices: List[int],
    old_text: str,
    new_text: str,
    mark_color: Optional[Tuple[int, int, int]] = (0, 128, 0),
) -> bool:
    """在候选 block 区域内改写。

    先逐段做精确/规范化替换；若原文跨多个段落，则在候选区域整体重写。
    """
    para_map = _ordered_paragraph_map(document)
    indices = [i for i in block_indices if i in para_map]
    for bi in indices:
        if apply_edit_at_block(document, bi, old_text, new_text, mark_color):
            return True
    if len(indices) <= 1:
        return False

    paras = [para_map[i] for i in indices]
    combined = "\n".join(para.text for para in paras)
    norm_old = _norm_text(old_text)
    norm_combined = _norm_text(combined)
    similar = difflib.SequenceMatcher(None, norm_old, norm_combined).ratio() if norm_old and norm_combined else 0.0
    if not (
        norm_combined and (
            norm_combined in norm_old
            or norm_old in norm_combined
            or similar >= 0.45
        )
    ):
        return False

    lines = _split_rewrite_lines(new_text, len(paras))
    for para, line in zip(paras, lines):
        _set_para_text(para, line, mark_color if line else None)
    return True


def locate_text_blocks(parsed: Dict[str, Any], needle: str,
                       hint_blocks: Optional[List[int]] = None) -> List[int]:
    """在段落模型中定位包含 needle 的 block_index（前端高亮兜底用）。
    优先在 hint_blocks（问题所属块的 source_blocks）内找，找不到再全局找。"""
    needle = (needle or "").strip()
    if not needle:
        return list(hint_blocks or [])
    blocks = parsed["blocks"]

    def search(indices):
        hits = [i for i in indices
                if 0 <= i < len(blocks) and needle in blocks[i]["text"]]
        if hits:
            return hits
        norm_needle = _norm_text(needle)
        return [i for i in indices
                if 0 <= i < len(blocks) and norm_needle and norm_needle in _norm_text(blocks[i]["text"])]

    if hint_blocks:
        hit = search(hint_blocks)
        if hit:
            return hit
        # 截断匹配（needle 可能跨段或含标点差异）
        short = needle[:12]
        hit = [i for i in hint_blocks
               if 0 <= i < len(blocks) and short and _norm_text(short) in _norm_text(blocks[i]["text"])]
        if hit:
            return hit
    return search(range(len(blocks)))


if __name__ == "__main__":
    import sys
    import tempfile
    # 自检：构造一个临时 docx，验证 block 索引 / 切块 / 改字 / 定位闭环
    doc = Document()
    doc.add_heading("第一章 安全技术措施", level=1)
    doc.add_paragraph("一、采煤机操作规程")
    doc.add_paragraph("3、起吊设备前，选择支护锚索作为起吊点。")
    doc.add_paragraph("4、所有被起吊设备均用4分钢丝绳套。")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "项目"
    table.cell(0, 1).text = "数值"
    table.cell(1, 0).text = "断电浓度"
    table.cell(1, 1).text = "1.2%"
    doc.add_paragraph("二、停风后的应急措施")
    doc.add_paragraph("恢复送风时间达5分钟后，上隅角瓦斯浓度不超过1.2%方可进入。")

    tmp = Path(tempfile.mkdtemp()) / "selftest.docx"
    doc.save(tmp)

    parsed, chunks = adapt(tmp)
    print(f"blocks: {parsed['n_blocks']}, chunks: {len(chunks)}")
    for c in chunks:
        print(f"  chunk source_blocks={c['source_blocks']} level={c['chunk_level']} "
              f"| {c['content'][:40].replace(chr(10),' ')}")

    # 定位：含"支护锚索"的段
    hits = locate_text_blocks(parsed, "支护锚索作为起吊点")
    print("locate 支护锚索 → blocks:", hits)
    assert hits, "应能定位到锚索段落"

    # 改字：把那一段的"支护锚索"改为"专用吊装锚索"
    d2 = Document(str(tmp))
    ok = apply_edit_at_block(d2, hits[0], "支护锚索", "专用吊装锚索")
    assert ok, "改字应成功"
    out = Path(tempfile.mkdtemp()) / "edited.docx"
    d2.save(out)
    reparsed = parse_docx(out)
    assert any("专用吊装锚索" in b["text"] for b in reparsed["blocks"]), "改字未生效"
    print("[OK] docx_adapter 自检通过：block索引/切块/表格/定位/改字闭环")
