"""Build reviewer-facing S1302 error-location artifacts from the blind benchmark.

Outputs are deliberately separate from the blind injected-error document:

1. a copy whose 30 injected errors are shown in red;
2. a compact Word catalog listing every error and its location.

Only ``word/document.xml`` is changed when producing the highlighted copy, so
the source document's sections, tables, fields, drawings, and embedded assets
remain untouched.
"""

from __future__ import annotations

import argparse
import copy
import json
import tempfile
import zipfile
from pathlib import Path
from typing import Any, Iterable

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor
from lxml import etree


PROJECT_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_DIR = (
    PROJECT_ROOT / "new_docs" / "test_doc" / "S1302人工错误注入测试集_v1"
)
BLIND_NAME = "004 S1302工作面作业规程（综采）_人工错误注入版_v1.docx"
GOLD_NAME = "004 S1302工作面作业规程（综采）_人工错误金标_v1.json"
HIGHLIGHTED_NAME = "004 S1302工作面作业规程（综采）_错误标红版_v1.docx"
CATALOG_NAME = "004 S1302工作面作业规程（综采）_错误位置清单_v1.docx"

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = {"w": W_NS}
XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"

# compact_reference_guide preset tokens.  The catalog uses a named landscape
# matrix override so five-column error tables remain readable.
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
HEADER_FILL = "E8EEF5"
ERROR_RED = "C00000"
TEXT = "202124"
MUTED = "5F6368"
WHITE = "FFFFFF"
TABLE_WIDTH_DXA = 15120
TABLE_INDENT_DXA = 120


def paragraph_text(paragraph: etree._Element) -> str:
    return "".join(
        node.text or "" for node in paragraph.xpath(".//w:t", namespaces=NS)
    )


def find_exact_paragraph(root: etree._Element, text: str) -> etree._Element:
    matches = [
        paragraph
        for paragraph in root.xpath(".//w:p", namespaces=NS)
        if paragraph_text(paragraph) == text
    ]
    if len(matches) != 1:
        raise ValueError(f"段落全文应唯一，实际 {len(matches)} 个: {text[:80]}")
    return matches[0]


def find_unique_paragraph(root: etree._Element, anchor: str) -> etree._Element:
    matches = [
        paragraph
        for paragraph in root.xpath(".//w:p", namespaces=NS)
        if anchor in paragraph_text(paragraph)
    ]
    if len(matches) != 1:
        raise ValueError(f"段落锚点应唯一，实际 {len(matches)} 个: {anchor}")
    return matches[0]


def set_text(node: etree._Element, value: str) -> None:
    node.text = value
    if value.startswith(" ") or value.endswith(" "):
        node.set(XML_SPACE, "preserve")
    else:
        node.attrib.pop(XML_SPACE, None)


def set_run_color(run: etree._Element, color: str) -> None:
    rpr = run.find(qn("w:rPr"))
    if rpr is None:
        rpr = etree.Element(qn("w:rPr"))
        run.insert(0, rpr)
    for existing in rpr.findall(qn("w:color")):
        rpr.remove(existing)
    color_node = etree.SubElement(rpr, qn("w:color"))
    color_node.set(qn("w:val"), color)


def set_paragraph_mark_color(paragraph: etree._Element, color: str) -> None:
    ppr = paragraph.find(qn("w:pPr"))
    if ppr is None:
        ppr = etree.Element(qn("w:pPr"))
        paragraph.insert(0, ppr)
    rpr = ppr.find(qn("w:rPr"))
    if rpr is None:
        rpr = etree.SubElement(ppr, qn("w:rPr"))
    for existing in rpr.findall(qn("w:color")):
        rpr.remove(existing)
    color_node = etree.SubElement(rpr, qn("w:color"))
    color_node.set(qn("w:val"), color)


def split_and_color_text_node(
    node: etree._Element, local_start: int, local_end: int, color: str
) -> None:
    value = node.text or ""
    if not 0 <= local_start < local_end <= len(value):
        raise ValueError("文本节点切分范围非法")

    run = node.getparent()
    if run is None or run.tag != qn("w:r"):
        raise ValueError("目标 w:t 不是普通 w:r 的直接子节点")
    run_text_nodes = run.findall(qn("w:t"))
    if len(run_text_nodes) != 1:
        raise ValueError("目标 run 含多个 w:t，拒绝扩大标红范围")

    parent = run.getparent()
    if parent is None:
        raise ValueError("目标 run 缺少父节点")
    index = parent.index(run)
    pieces = [
        (value[:local_start], False),
        (value[local_start:local_end], True),
        (value[local_end:], False),
    ]
    inserted = 0
    for segment, is_error in pieces:
        if not segment:
            continue
        clone = copy.deepcopy(run)
        clone_text_nodes = clone.findall(qn("w:t"))
        if len(clone_text_nodes) != 1:
            raise ValueError("复制 run 后文本节点数量异常")
        set_text(clone_text_nodes[0], segment)
        if is_error:
            set_run_color(clone, color)
        parent.insert(index + inserted, clone)
        inserted += 1
    parent.remove(run)


def color_substring(paragraph: etree._Element, target: str) -> None:
    text_nodes = paragraph.xpath(".//w:t", namespaces=NS)
    full_text = "".join(node.text or "" for node in text_nodes)
    start = full_text.find(target)
    if start < 0:
        raise ValueError(f"段落内未找到待标红文本: {target}")
    if full_text.find(target, start + 1) >= 0:
        raise ValueError(f"段落内待标红文本不唯一: {target}")
    end = start + len(target)

    spans: list[tuple[etree._Element, int, int]] = []
    cursor = 0
    for node in text_nodes:
        value = node.text or ""
        spans.append((node, cursor, cursor + len(value)))
        cursor += len(value)

    # Reverse order keeps the precomputed offsets valid while runs are split.
    for node, node_start, node_end in reversed(spans):
        overlap_start = max(start, node_start)
        overlap_end = min(end, node_end)
        if overlap_start >= overlap_end:
            continue
        split_and_color_text_node(
            node,
            overlap_start - node_start,
            overlap_end - node_start,
            ERROR_RED,
        )


def next_paragraph_sibling(paragraph: etree._Element) -> etree._Element:
    sibling = paragraph.getnext()
    while sibling is not None and sibling.tag != qn("w:p"):
        sibling = sibling.getnext()
    if sibling is None:
        raise ValueError("插入位置后未找到重复段落")
    return sibling


def find_destination_with_duplicate(
    root: etree._Element, anchor: str, duplicated_text: str
) -> tuple[etree._Element, etree._Element]:
    """Resolve a destination even when repetition injections made its anchor non-unique."""
    candidates = [
        paragraph
        for paragraph in root.xpath(".//w:p", namespaces=NS)
        if anchor in paragraph_text(paragraph)
    ]
    matched: list[tuple[etree._Element, etree._Element]] = []
    for candidate in candidates:
        try:
            following = next_paragraph_sibling(candidate)
        except ValueError:
            continue
        if paragraph_text(following) == duplicated_text:
            matched.append((candidate, following))
    if len(matched) != 1:
        raise ValueError(
            f"插入位置与重复段落组合应唯一，实际 {len(matched)} 个: {anchor}"
        )
    return matched[0]


def color_entire_paragraph(paragraph: etree._Element) -> None:
    runs = paragraph.xpath(".//w:r", namespaces=NS)
    if not runs:
        raise ValueError("重复段落不含可标红的 run")
    for run in runs:
        set_run_color(run, ERROR_RED)
    set_paragraph_mark_color(paragraph, ERROR_RED)


def build_highlighted_docx(
    blind_path: Path, gold: dict[str, Any], output_path: Path
) -> dict[str, int]:
    with zipfile.ZipFile(blind_path, "r") as source_zip:
        document_xml = source_zip.read("word/document.xml")
        package_items = [(item, source_zip.read(item.filename)) for item in source_zip.infolist()]

    parser = etree.XMLParser(remove_blank_text=False, resolve_entities=False)
    root = etree.fromstring(document_xml, parser)
    inline_count = 0
    repetition_count = 0

    for case in gold["cases"]:
        if case["type"] in {"compliance", "typo"}:
            paragraph = find_exact_paragraph(root, case["mutated_paragraph"])
            color_substring(paragraph, case["mutated_text"])
            inline_count += 1
            continue

        if case["type"] == "redundancy":
            _destination, inserted = find_destination_with_duplicate(
                root, case["destination_anchor"], case["duplicated_text"]
            )
            actual = paragraph_text(inserted)
            if actual != case["duplicated_text"]:
                raise ValueError(
                    f"{case['error_id']} 重复段落不匹配: {actual[:80]}"
                )
            color_entire_paragraph(inserted)
            repetition_count += 1
            continue

        raise ValueError(f"未知错误类型: {case['type']}")

    highlighted_xml = etree.tostring(
        root, xml_declaration=True, encoding="UTF-8", standalone="yes"
    )
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(
        suffix=".docx", dir=output_path.parent, delete=False
    ) as temporary:
        temporary_path = Path(temporary.name)
    try:
        with zipfile.ZipFile(temporary_path, "w") as target_zip:
            for item, data in package_items:
                if item.filename == "word/document.xml":
                    data = highlighted_xml
                target_zip.writestr(item, data)
        temporary_path.replace(output_path)
    finally:
        if temporary_path.exists():
            temporary_path.unlink()

    return {
        "inline_red": inline_count,
        "repetition_red": repetition_count,
        "total_red_cases": inline_count + repetition_count,
    }


def set_cell_shading(cell, fill: str) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcpr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 80, bottom: int = 80, start: int = 120, end: int = 120) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    tc_mar = tcpr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tcpr.append(tc_mar)
    for edge, value in (("top", top), ("bottom", bottom), ("start", start), ("end", end)):
        element = tc_mar.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            tc_mar.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa: int) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    tcw = tcpr.find(qn("w:tcW"))
    if tcw is None:
        tcw = OxmlElement("w:tcW")
        tcpr.append(tcw)
    tcw.set(qn("w:w"), str(width_dxa))
    tcw.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: list[int]) -> None:
    if sum(widths) != TABLE_WIDTH_DXA:
        raise ValueError(f"表格列宽总和必须为 {TABLE_WIDTH_DXA}: {widths}")
    table.autofit = False
    tblpr = table._tbl.tblPr
    tblw = tblpr.find(qn("w:tblW"))
    if tblw is None:
        tblw = OxmlElement("w:tblW")
        tblpr.append(tblw)
    tblw.set(qn("w:w"), str(TABLE_WIDTH_DXA))
    tblw.set(qn("w:type"), "dxa")

    tbl_ind = tblpr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tblpr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tblpr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tblpr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for row in table.rows:
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        trpr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        trpr.append(cant_split)
        for cell, width in zip(row.cells, widths):
            set_cell_width(cell, width)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_table_header(row) -> None:
    trpr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    trpr.append(header)


def set_run_font(
    run,
    *,
    size: float = 9,
    color: str = TEXT,
    bold: bool = False,
    italic: bool = False,
) -> None:
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic


def style_paragraph(paragraph, *, after: float = 2, line: float = 1.1) -> None:
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line
    paragraph.paragraph_format.widow_control = True


def clear_cell(cell) -> None:
    paragraph = cell.paragraphs[0]
    paragraph.clear()
    style_paragraph(paragraph, after=0, line=1.05)


def add_cell_text(
    cell,
    parts: Iterable[tuple[str, str, bool]],
    *,
    size: float = 8.4,
    align=WD_ALIGN_PARAGRAPH.LEFT,
) -> None:
    clear_cell(cell)
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    for text, color, bold in parts:
        run = paragraph.add_run(text)
        set_run_font(run, size=size, color=color, bold=bold)


def add_page_field(paragraph, field: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = f" {field} "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, text, end])
    set_run_font(run, size=8.5, color=MUTED)


def configure_styles(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    normal.font.size = Pt(11)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    heading_specs = {
        "Title": (23, TEXT, 0, 4),
        "Subtitle": (12, MUTED, 0, 14),
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 14, 7),
        "Heading 3": (12, DARK_BLUE, 10, 5),
    }
    for style_name, (size, color, before, after) in heading_specs.items():
        style = document.styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = style_name != "Subtitle"
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def configure_section(document: Document) -> None:
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Mm(297)
    section.page_height = Mm(210)
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)
    section.header_distance = Inches(0.24)
    section.footer_distance = Inches(0.24)

    header = section.header
    paragraph = header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    style_paragraph(paragraph, after=0, line=1.0)
    left = paragraph.add_run("S1302 人工错误注入测试集")
    set_run_font(left, size=8.5, color=MUTED, bold=True)
    right = paragraph.add_run("  |  错误位置清单 v1")
    set_run_font(right, size=8.5, color=MUTED)

    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    style_paragraph(paragraph, after=0, line=1.0)
    prefix = paragraph.add_run("第 ")
    set_run_font(prefix, size=8.5, color=MUTED)
    add_page_field(paragraph, "PAGE")
    middle = paragraph.add_run(" 页 / 共 ")
    set_run_font(middle, size=8.5, color=MUTED)
    add_page_field(paragraph, "NUMPAGES")
    suffix = paragraph.add_run(" 页")
    set_run_font(suffix, size=8.5, color=MUTED)


def add_bottom_rule(paragraph, color: str = BLUE, size: int = 12) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    borders = ppr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        ppr.append(borders)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "5")
    bottom.set(qn("w:color"), color)
    borders.append(bottom)


def add_masthead(document: Document, counts: dict[str, int]) -> None:
    kicker = document.add_paragraph()
    style_paragraph(kicker, after=3, line=1.0)
    run = kicker.add_run("人工构造错误 · 核对专用")
    set_run_font(run, size=9.5, color=BLUE, bold=True)

    title = document.add_paragraph(style="Title")
    title.add_run("S1302 工作面作业规程（综采）错误位置清单")
    subtitle = document.add_paragraph(style="Subtitle")
    subtitle.add_run("用于核验多智能体对合规性错误、错别字和跨章节重复内容的识别情况")

    metadata = document.add_table(rows=3, cols=2)
    metadata.style = "Table Grid"
    metadata_rows = [
        ("基准文档", BLIND_NAME),
        ("金标文件", GOLD_NAME),
        ("使用说明", "盲测时不要向系统提供本清单或标红版；系统输出完成后再据此核对。"),
    ]
    for row, (label, value) in zip(metadata.rows, metadata_rows):
        set_cell_shading(row.cells[0], HEADER_FILL)
        add_cell_text(row.cells[0], [(label, DARK_BLUE, True)], size=8.8)
        add_cell_text(row.cells[1], [(value, TEXT, False)], size=8.8)
    set_table_geometry(metadata, [2160, 12960])

    spacer = document.add_paragraph()
    style_paragraph(spacer, after=2, line=1.0)

    metric = document.add_table(rows=1, cols=4)
    metric.style = "Table Grid"
    metrics = [
        ("合规性错误", counts["compliance"]),
        ("错别字错误", counts["typo"]),
        ("重复内容", counts["redundancy"]),
        ("合计", counts["total"]),
    ]
    for cell, (label, value) in zip(metric.rows[0].cells, metrics):
        set_cell_shading(cell, HEADER_FILL)
        add_cell_text(
            cell,
            [(f"{value}", BLUE, True), (f"\n{label}", MUTED, False)],
            size=10.5,
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )
    set_table_geometry(metric, [3780, 3780, 3780, 3780])

    note = document.add_paragraph()
    style_paragraph(note, after=4, line=1.15)
    run = note.add_run("标记规则：")
    set_run_font(run, size=9.5, color=TEXT, bold=True)
    run = note.add_run("红色文字")
    set_run_font(run, size=9.5, color=ERROR_RED, bold=True)
    run = note.add_run("为人为注入的错误；重复类在标红版中将整段插入副本标红，原始段落保持原色。")
    set_run_font(run, size=9.5, color=TEXT)
    add_bottom_rule(note)


def add_header_row(table, labels: list[str]) -> None:
    row = table.rows[0]
    set_repeat_table_header(row)
    for cell, label in zip(row.cells, labels):
        set_cell_shading(cell, HEADER_FILL)
        add_cell_text(
            cell,
            [(label, DARK_BLUE, True)],
            size=8.4,
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )


def short_context(text: str, target: str, radius: int = 42) -> str:
    index = text.find(target)
    if index < 0:
        return text[: radius * 2] + ("…" if len(text) > radius * 2 else "")
    start = max(0, index - radius)
    end = min(len(text), index + len(target) + radius)
    prefix = "…" if start else ""
    suffix = "…" if end < len(text) else ""
    return prefix + text[start:end] + suffix


def add_compliance_table(document: Document, cases: list[dict[str, Any]]) -> None:
    document.add_heading("一、合规性错误（12 项）", level=2)
    intro = document.add_paragraph()
    style_paragraph(intro, after=5, line=1.15)
    run = intro.add_run("核对重点：阈值放宽、停工/停电要求取消、允许带电检修等安全约束变化。")
    set_run_font(run, size=9.2, color=MUTED)

    table = document.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    add_header_row(table, ["编号", "原块", "正确原文", "注入错误", "判定依据"])
    for case in cases:
        cells = table.add_row().cells
        add_cell_text(cells[0], [(case["error_id"], BLUE, True)], align=WD_ALIGN_PARAGRAPH.CENTER)
        add_cell_text(cells[1], [(str(case["baseline_block_index"]), TEXT, False)], align=WD_ALIGN_PARAGRAPH.CENTER)
        add_cell_text(cells[2], [(case["original_text"], TEXT, False)])
        add_cell_text(cells[3], [(case["mutated_text"], ERROR_RED, True)])
        add_cell_text(cells[4], [(case["basis"], TEXT, False)])
    set_table_geometry(table, [720, 840, 3360, 3360, 6840])


def add_typo_table(document: Document, cases: list[dict[str, Any]]) -> None:
    document.add_page_break()
    document.add_heading("二、错别字错误（10 项）", level=2)
    intro = document.add_paragraph()
    style_paragraph(intro, after=5, line=1.15)
    run = intro.add_run("核对重点：专业术语中的同音字、形近字和单字替换。")
    set_run_font(run, size=9.2, color=MUTED)

    table = document.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    add_header_row(table, ["编号", "原块", "错误字词", "正确字词", "所在段落摘录"])
    for case in cases:
        cells = table.add_row().cells
        add_cell_text(cells[0], [(case["error_id"], BLUE, True)], align=WD_ALIGN_PARAGRAPH.CENTER)
        add_cell_text(cells[1], [(str(case["baseline_block_index"]), TEXT, False)], align=WD_ALIGN_PARAGRAPH.CENTER)
        add_cell_text(cells[2], [(case["mutated_text"], ERROR_RED, True)])
        add_cell_text(cells[3], [(case["original_text"], TEXT, False)])
        context = short_context(case["mutated_paragraph"], case["mutated_text"])
        add_cell_text(cells[4], [(context, TEXT, False)])
    set_table_geometry(table, [720, 840, 2160, 2160, 9240])


def add_repetition_table(document: Document, cases: list[dict[str, Any]]) -> None:
    document.add_page_break()
    document.add_heading("三、重复内容错误（8 项）", level=2)
    intro = document.add_paragraph()
    style_paragraph(intro, after=5, line=1.15)
    run = intro.add_run("核对重点：下列段落从原位置完整复制，并插入到另一章节；标红版只标红插入副本。")
    set_run_font(run, size=9.2, color=MUTED)

    table = document.add_table(rows=1, cols=6)
    table.style = "Table Grid"
    add_header_row(table, ["编号", "源块", "插入于块后", "源位置锚点", "插入位置锚点", "重复段落内容"])
    for case in cases:
        cells = table.add_row().cells
        add_cell_text(cells[0], [(case["error_id"], BLUE, True)], align=WD_ALIGN_PARAGRAPH.CENTER)
        add_cell_text(cells[1], [(str(case["source_block_index"]), TEXT, False)], align=WD_ALIGN_PARAGRAPH.CENTER)
        add_cell_text(cells[2], [(str(case["inserted_after_block_index"]), TEXT, False)], align=WD_ALIGN_PARAGRAPH.CENTER)
        add_cell_text(cells[3], [(case["source_anchor"], TEXT, False)])
        add_cell_text(cells[4], [(case["destination_anchor"], TEXT, False)])
        duplicated = case["duplicated_text"]
        excerpt = duplicated[:180] + ("…" if len(duplicated) > 180 else "")
        add_cell_text(cells[5], [(excerpt, ERROR_RED, True)])
    set_table_geometry(table, [660, 780, 1020, 2880, 2880, 6900])


def build_catalog(gold: dict[str, Any], output_path: Path) -> None:
    document = Document()
    document.core_properties.title = "S1302工作面作业规程（综采）错误位置清单"
    document.core_properties.subject = "人工错误注入测试集金标核对"
    document.core_properties.author = "Coal Regulation AI Evaluation"
    document.core_properties.keywords = "S1302, 合规审查, 错别字, 重复内容, 金标"
    configure_styles(document)
    configure_section(document)
    add_masthead(document, gold["counts"])

    cases = gold["cases"]
    add_compliance_table(document, [case for case in cases if case["type"] == "compliance"])
    add_typo_table(document, [case for case in cases if case["type"] == "typo"])
    add_repetition_table(document, [case for case in cases if case["type"] == "redundancy"])

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)


def validate_outputs(
    blind_path: Path,
    highlighted_path: Path,
    catalog_path: Path,
    gold: dict[str, Any],
) -> dict[str, Any]:
    with zipfile.ZipFile(blind_path) as blind_zip, zipfile.ZipFile(highlighted_path) as red_zip:
        blind_names = set(blind_zip.namelist())
        red_names = set(red_zip.namelist())
        if blind_names != red_names:
            raise ValueError("标红版 DOCX 包结构与盲测版不一致")
        changed_parts = [
            name
            for name in sorted(blind_names)
            if blind_zip.read(name) != red_zip.read(name)
        ]
        if changed_parts != ["word/document.xml"]:
            raise ValueError(f"标红版出现非预期包修改: {changed_parts}")

        red_root = etree.fromstring(red_zip.read("word/document.xml"))
        red_runs = red_root.xpath(
            ".//w:r[w:rPr/w:color[@w:val='C00000']]", namespaces=NS
        )
        if not red_runs:
            raise ValueError("标红版未发现红色 run")

    catalog = Document(catalog_path)
    if len(catalog.tables) != 5:
        raise ValueError(f"错误清单表格数量异常: {len(catalog.tables)}")
    catalog_text = "\n".join(p.text for p in catalog.paragraphs)
    for expected in ("合规性错误", "错别字错误", "重复内容错误"):
        if expected not in catalog_text:
            raise ValueError(f"错误清单缺少章节: {expected}")

    return {
        "changed_parts": changed_parts,
        "red_run_count": len(red_runs),
        "catalog_tables": len(catalog.tables),
        "catalog_sections": len(catalog.sections),
        "case_count": len(gold["cases"]),
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", type=Path, default=DEFAULT_DIR)
    args = parser.parse_args()

    input_dir = args.input_dir.resolve()
    blind_path = input_dir / BLIND_NAME
    gold_path = input_dir / GOLD_NAME
    highlighted_path = input_dir / HIGHLIGHTED_NAME
    catalog_path = input_dir / CATALOG_NAME
    if not blind_path.exists():
        raise FileNotFoundError(blind_path)
    if not gold_path.exists():
        raise FileNotFoundError(gold_path)

    gold = json.loads(gold_path.read_text(encoding="utf-8"))
    if gold["counts"]["total"] != 30 or len(gold["cases"]) != 30:
        raise ValueError("金标错误数量不是预期的 30 项")

    highlighted_counts = build_highlighted_docx(blind_path, gold, highlighted_path)
    build_catalog(gold, catalog_path)
    validation = validate_outputs(blind_path, highlighted_path, catalog_path, gold)
    print(
        json.dumps(
            {
                "highlighted_docx": str(highlighted_path),
                "catalog_docx": str(catalog_path),
                "highlighted_counts": highlighted_counts,
                "validation": validation,
            },
            ensure_ascii=False,
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
