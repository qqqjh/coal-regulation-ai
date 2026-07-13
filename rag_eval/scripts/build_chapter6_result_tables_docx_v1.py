# -*- coding: utf-8 -*-
"""生成一个独立的 docx：第六章 6.1-6.4 四张结果表（数据取自评测 JSON）。

风格与论文设计表一致：表头深蓝底(1F4E79)+白色粗体，数据深色(222222)，Table Grid 边框，宋体 10.5pt。
"""
import json
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

for _s in (sys.stdout, sys.stderr):
    try:
        _s.reconfigure(encoding="utf-8")
    except Exception:
        pass

ROOT = Path(__file__).resolve().parents[2]
OUT = Path(r"D:/work/项目/研电赛/第六章6.1-6.4结果表.docx")
CH6 = json.loads((ROOT / "rag_eval" / "data" / "chapter6_evaluation_report_v1.json").read_text(encoding="utf-8"))

abl = CH6["retrieval_ablation"]["variants"]
sg = CH6["section_6_3_strategy_gen_eval"]
par = {r["concurrency"]: r for r in CH6["parallel_performance_eval"]["runs"]}
gnd = CH6["groundedness"]
func = CH6["section_6_2_function_quality"]
agents = CH6["agent_output_summary"]
ragas = (CH6.get("ragas_metrics") or {}).get("metrics") or {}

HEADER_FILL = "1F4E79"
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK = RGBColor(0x22, 0x22, 0x22)
EA_FONT = "宋体"


def f4(x):
    return "-" if x is None else f"{x:.4f}"


def pc(x):
    return "-" if x is None else f"{x * 100:.2f}%"


def f1(x):
    return "-" if x is None else f"{x:.1f}"


def f2(x):
    return "-" if x is None else f"{x:.2f}"


def gm(block, *names):
    m = (gnd.get(block) or {}).get("metrics") or {}
    for n in names:
        row = m.get(n)
        if isinstance(row, dict) and row.get("mean") is not None:
            return row["mean"]
    return None


def rm(name):
    row = ragas.get(name)
    return row.get("mean") if isinstance(row, dict) else None


pm = abl["parent_multi"]["aggregate"]

# -------- 四张表内容 --------
T61 = [["检索方案", "覆盖率(Hit@15)", "Hit@5", "Recall@5", "MRR", "NDCG@10", "平均耗时/ms"]]
for key, name in [("dense", "稠密检索基线"), ("hybrid", "混合检索"),
                  ("hybrid_rerank", "混合检索+重排"), ("parent_multi", "父块保底+片段补充")]:
    a = abl[key]["aggregate"]
    T61.append([name, f4(a["hit_at"]["15"]), f4(a["hit_at"]["5"]), f4(a["recall_at"]["5"]),
                f4(a["mrr"]), f4(a["ndcg_at"]["10"]), f1(abl[key]["avg_retrieval_ms"])])

T62 = [["测试对象", "评测指标", "结果", "数据规模 / 说明"],
       ["知识检索智能体", "Hit@5 / Recall@5 / MRR",
        f"{f4(pm['hit_at']['5'])} / {f4(pm['recall_at']['5'])} / {f4(pm['mrr'])}",
        "父块保底+片段补充，44 个可评价案例（详见表 6-1）"],
       ["合规审查智能体", "结论匹配准确率 / 非合规 F1",
        f"{pc(func['exact_status_accuracy'])} / {pc(func['non_compliance_f1'])}",
        f"人工黄金结论匹配；审查链 RAGAS Faithfulness {f4(rm('faithfulness'))}、"
        f"Relevancy {f4(rm('answer_relevancy') or rm('response_relevancy'))}"],
       ["错别字检查智能体", "Faithfulness / ResponseRelevancy",
        f"{f4(gm('typo', 'faithfulness'))} / {f4(gm('typo', 'answer_relevancy', 'response_relevancy'))}",
        f"无人工金标，RAGAS 有据性代理，{agents['typo_agent']['chunks_with_typo_issues']} 块"
        "（纠错为推断性结论，Faithfulness 偏保守）"],
       ["重复性检查智能体", "Faithfulness / ResponseRelevancy",
        f"{f4(gm('redundancy', 'faithfulness'))} / {f4(gm('redundancy', 'answer_relevancy', 'response_relevancy'))}",
        f"无人工金标，RAGAS 有据性代理，{agents['repetition_agent']['chunks_with_duplicates']} 块"],
       ["数值核验智能体", "触发覆盖",
        f"{agents['numeric_agent']['chunks_with_numeric_checks']} 块触发确定性数值核验",
        "单位归一化+方向语义比较；准确率/漏报率待对照测试"],
       ["结果修订智能体", "定位与修改", "端到端闭环已验证", "批量定位/修改成功率待测"],
       ["现场问答智能体", "—", "规划中，尚未实现", "策略可规划（见表 6-3），执行待实现"]]

T63 = [["任务类型", "平均生成时间/s", "任务识别", "工具选择F1", "策略执行匹配", "升级决策"]]
for t in ["规则文档入库", "待审文档审查", "争议项复核", "Word修订", "现场风险问答"]:
    r = sg["by_type"][t]
    T63.append([t, f2(r["mean_elapsed_sec"]), pc(r["task_recognition_accuracy"]),
                pc(r["mean_tool_f1"]), pc(r["mean_step_match"]), pc(r["escalation_decision_accuracy"])])
ag = sg["aggregate"]
T63.append(["全部任务", f2(ag["mean_elapsed_sec"]), pc(ag["task_recognition_accuracy"]),
            pc(ag["mean_tool_f1"]), pc(ag["mean_step_match"]), pc(ag["escalation_decision_accuracy"])])

r1, r4 = par[1], par[4]
T64 = [["执行模式", "块数", "总耗时/s", "首条返回/s", "吞吐(块/分)", "加速比", "异常率", "结果一致率"],
       ["串行基线（并发=1）", str(r1["chunk_count"]), f1(r1["total_elapsed_sec"]),
        f1(r1["first_issue_latency_sec"]), f2(r1["throughput_chunks_per_min"]),
        f"{r1['speedup_vs_serial']:.2f}×", pc(r1["error_rate"]), "基准"],
       ["链并行+块级并发（并发=4）", str(r4["chunk_count"]), f1(r4["total_elapsed_sec"]),
        f1(r4["first_issue_latency_sec"]), f2(r4["throughput_chunks_per_min"]),
        f"{r4['speedup_vs_serial']:.2f}×", pc(r4["error_rate"]), pc(r4["result_consistency"])]]

SECTIONS = [
    ("表 6-1  四档检索方案效果对比（统一黄金集 "
     f"{CH6['retrieval_ablation']['evaluated_case_count']} 个可评价案例，TopK={CH6['retrieval_ablation']['topk']}；"
     "重排两档关阈值做纯排序消融）", T61),
    ("表 6-2  各专业智能体功能测试结果", T62),
    (f"表 6-3  主智能体策略生成评测结果（{ag['count']} 个任务覆盖五类，P95={ag['p95_elapsed_sec']:.2f}s）", T63),
    ("表 6-4  串行与并行审查性能对比（文档 066，142 块；合规/错别字/重复三链恒并行，"
     "并发数控制块级并发，并发=1 即链并行块串行）", T64),
]


# -------- 构建文档 --------
def set_run(run, text, color, bold, size=10.5):
    run.text = text
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.size = Pt(size)
    run.font.name = "Times New Roman"
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), EA_FONT)


def set_shd(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:fill'), fill)
    tcPr.append(shd)


doc = Document()
doc.styles['Normal'].font.name = "Times New Roman"
doc.styles['Normal'].font.size = Pt(10.5)
doc.styles['Normal'].element.rPr.rFonts.set(qn('w:eastAsia'), EA_FONT)

title = doc.add_paragraph()
tr = title.add_run("第六章 6.1–6.4 系统测试结果表")
tr.font.bold = True
tr.font.size = Pt(14)
tr.font.name = "Times New Roman"
tr._element.get_or_add_rPr().append(OxmlElement('w:rFonts'))
title.runs[0]._element.rPr.rFonts.set(qn('w:eastAsia'), "黑体")
note = doc.add_paragraph()
nr = note.add_run(f"数据来源：{CH6.get('generated_at', '')} 评测产物；详见 rag_eval/reports/*。")
nr.font.size = Pt(9)
nr.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
nr.font.name = "Times New Roman"
nr._element.get_or_add_rPr().append(OxmlElement('w:rFonts'))
nr._element.rPr.rFonts.set(qn('w:eastAsia'), EA_FONT)

for caption, data in SECTIONS:
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.LEFT
    cr = cap.add_run(caption)
    cr.font.bold = True
    cr.font.size = Pt(10.5)
    cr.font.color.rgb = DARK
    cr.font.name = "Times New Roman"
    cr._element.get_or_add_rPr().append(OxmlElement('w:rFonts'))
    cap.runs[0]._element.rPr.rFonts.set(qn('w:eastAsia'), EA_FONT)

    nrow, ncol = len(data), len(data[0])
    table = doc.add_table(rows=nrow, cols=ncol)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for ri, row in enumerate(data):
        for ci, val in enumerate(row):
            cell = table.rows[ri].cells[ci]
            if ri == 0:
                set_shd(cell, HEADER_FILL)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if (ci > 0 or ri == 0) else WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run()
            set_run(run, val, WHITE if ri == 0 else DARK, ri == 0)
    doc.add_paragraph()  # 表间空行

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(str(OUT))
print(f"已生成: {OUT}")
print(f"表数: {len(doc.tables)}（应为 4）")
